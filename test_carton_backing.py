"""
Carton-backing tests -- which quantity rows are shipment data.

Offline and free: real workbooks from the repo plus synthetic ones built here.
Pins Paula's ruling of 2026-09-23 (a quantity row counts only if cartons in the
same document back it) and, as importantly, that the rule reads ARITHMETIC and
never a row's label.

    python test_carton_backing.py
"""

from __future__ import annotations

import tempfile
import traceback
from pathlib import Path

import carton_backing as cb
from claude_extractor import read_workbook_grids

HERE = Path(__file__).resolve().parent
FOOTWEAR = HERE / "FW26 footwear PO-1624 packing sheets (invoice sheet removed).xlsx"
TAINAN = HERE / "50144--- PO 0001725   packing list  ( Correction on Aug.10 from Aug.07 ).xls"
_results: list[tuple[bool, str, str]] = []


def check(ok: bool, name: str, detail: str = "") -> bool:
    _results.append((bool(ok), name, detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  --  {detail}" if detail else ""))
    return bool(ok)


def section(title: str) -> None:
    print()
    print(f"--- {title} " + "-" * max(0, 70 - len(title)))


def _sheet(path: Path, name: str) -> cb._Sheet:
    return cb._Sheet(next(g for g in read_workbook_grids(path) if g.name == name))


def footwear_like(path: Path, labels=("By Sea", "Ordered Qty", "By UPS"),
                  typo_row=False) -> Path:
    """
    PO 1624's layout: sizes 8..14 in K..Q, per-carton R, cartons S, total T.

    Three quantity rows under the cartons, labelled by `labels` -- in order, the
    carton-backed row, the order row, and the later-shipment row. Swapping the
    labels is how the tests prove nothing reads them.
    """
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "PO-9999 20138"
    header = [""] * 10 + [8, 9, 10, 11, 12, 13, 14]
    sheet.append(["STYLE"] + [""] * 9 + header[10:])        # row 1: size header
    cartons = [  # (first, last, size index 0..6, qty per carton, cartons)
        (1, 3, 1, 8, 3), (4, 10, 2, 8, 7), (11, 19, 3, 8, 9), (20, 25, 4, 8, 6),
    ]
    for first, last, index, per, count in cartons:            # rows 2..5
        row = [""] * 20
        row[1], row[3] = first, last
        row[10 + index] = per * count
        row[17], row[18], row[19] = per, count, per * count
        sheet.append(row)
    shipped = [0, 24, 56, 72, 48, 0, 0]
    if typo_row:                                              # row 6: 511/533's typo
        row = [""] * 20
        row[1], row[3] = 26, 26
        row[15] = 20
        row[17], row[18], row[19] = 11, 1, 20   # "11 per carton" beside 20 pieces
        sheet.append(row)
        shipped[5] = 20
    order = [10, 40, 75, 85, 55, 25, 10]
    later = [o - s for o, s in zip(order, shipped)]
    for label, values in zip(labels, (shipped, order, later)):
        row = [""] * 20
        row[9] = label
        for i, v in enumerate(values):
            row[10 + i] = v or ""
        sheet.append(row)
    book.save(path)
    return path


def rows_of(path: Path) -> dict[str, int]:
    """Row number of each quantity row, found by position (they are the last three)."""
    grid = read_workbook_grids(path)[0]
    n = grid.n_rows
    return {"shipped": n - 2, "order": n - 1, "later": n}


# ---------------------------------------------------------------------------


def test_po_1624_real_sheets() -> None:
    section("PO 1624, the real footwear sheets: only the carton-backed row is shipment data")
    s20138 = _sheet(FOOTWEAR, "PO-1624 20138")
    check(s20138.backing(39).status == cb.BACKED,
          "20138 R39 (printed 'By Sea') = the sum of carton rows 29-38",
          s20138.backing(39).reason)
    check(s20138.backing(40).status == cb.UNBACKED,
          "20138 R40 (printed 'Ordered Qty') is not backed by any cartons")
    check(s20138.backing(41).status == cb.UNBACKED,
          "20138 R41 (printed 'By UPS') is not backed -- the later shipment Paula ruled on")

    s20139 = _sheet(FOOTWEAR, "PO-1624 20139")
    v56 = s20139.backing(56)
    check(v56.status == cb.BACKED and v56.carton_rows[0] == 46,
          "20139 MLT R56 is backed by ITS OWN cartons (46-55), not DKF's identical ones",
          v56.reason)
    check(s20139.backing(57).status == cb.UNBACKED and s20139.backing(58).status == cb.UNBACKED,
          "20139 MLT's order and later rows are unbacked even though DKF's rows are "
          "numerically identical to them -- a summary row never counts as a carton")

    v = _sheet(FOOTWEAR, "PO-1624 20140").backing(39)
    check(v.status == cb.BACKED and v.carton_rows == tuple(range(29, 39)),
          "20140 R39, the UNLABELLED row, is backed by carton rows 29-38 -- a full "
          "shipment, kept", v.reason)


#: PO 1624's sea shipment, from the stored HBL DRAFT.pdf classifier rationale
#: ("255 cartons / 2040 pairs"). A FIXED expected value -- the HBL is not read as a
#: data source here or anywhere; it is the independent figure the slip must meet.
HBL_SEA_PAIRS = 2040
#: What the slip prints as going later by UPS (Paula, 2026-09-23: not this shipment).
UPS_LATER_PAIRS = 60


def test_po_1624_sea_shipment_matches_the_hbl() -> None:
    section("PO 1624: the carton-backed rows ARE the sea shipment -- 2,040 pairs, per the HBL")
    sheets = {g.name: cb._Sheet(g) for g in read_workbook_grids(FOOTWEAR)}

    def total(rows):
        return sum(sum(sheets[s].sizes(r).values()) for s, r in rows)

    sea = [("PO-1624 20138", 39), ("PO-1624 20139", 39), ("PO-1624 20139", 56),
           ("PO-1624 20140", 39)]  # three 'By Sea' rows + 20140's unlabelled row
    ups = [("PO-1624 20138", 41), ("PO-1624 20139", 41), ("PO-1624 20139", 58)]
    check(total(sea) == HBL_SEA_PAIRS,
          "By Sea rows + the unlabelled 20140 row = 2,040 pairs, the HBL's figure",
          f"{total(sea):g}")
    check(total(ups) == UPS_LATER_PAIRS, "By UPS rows = 60 pairs", f"{total(ups):g}")
    check(all(sheets[s].backing(r).status == cb.BACKED for s, r in sea)
          and all(sheets[s].backing(r).status == cb.UNBACKED for s, r in ups),
          "and the carton rule keeps exactly the 2,040 and removes exactly the 60")
    backed_quantity_rows = sorted(
        (name, r) for name, sheet in sheets.items() for r in range(1, sheet.grid.n_rows + 1)
        if r not in {c for c, _ in sheet.cartons} and sheet.backing(r).status == cb.BACKED)
    check(backed_quantity_rows == sorted(sea),
          "no OTHER summary row on the workbook is backed -- the sea rows are the whole "
          "shipment", str(backed_quantity_rows))
    cartons = sum(sum(c.values()) for sheet in sheets.values() for _r, c in sheet.cartons)
    check(cartons == HBL_SEA_PAIRS,
          "independently: the carton rows themselves total 2,040", f"{cartons:g}")


def test_po_1725_sheets_are_both_backed() -> None:
    section("PO 1725: ACTUAL rows backed on BOTH sheets, ORDER rows on neither")
    for name, actual, order in (("ACT", (47, 53, 121, 126), (46, 52, 120, 125)),
                                ("REV", (47, 53, 121, 127), (46, 52, 120, 126))):
        sheet = _sheet(TAINAN, name)
        check(all(sheet.backing(r).status == cb.BACKED for r in actual),
              f"{name}: every ACTUAL row is carton-backed (per-carton x count layout)")
        check(all(sheet.backing(r).status == cb.UNBACKED for r in order),
              f"{name}: every ORDER row is not")
    # So this rule does NOT choose between ACT and REV -- each sheet has its OWN
    # carton rows (same row numbers, 22 of 26 differ), summing to 865 and 860.
    # What keeps REV out is that the vendor HID it: `claude_extractor.visible_grids`.


def test_labels_are_never_read() -> None:
    section("the rule reads arithmetic, not labels: swap the labels, same verdicts")
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        plain = footwear_like(tmp / "plain.xlsx")
        swapped = footwear_like(tmp / "swapped.xlsx",
                                labels=("By UPS", "By Sea", "ACTUAL SHIPPED"))
        blank = footwear_like(tmp / "blank.xlsx", labels=("", "", ""))
        for path, what in ((plain, "vendor labels"), (swapped, "misleading labels"),
                           (blank, "no labels at all")):
            sheet = cb._Sheet(read_workbook_grids(path)[0])
            rows = rows_of(path)
            got = {k: sheet.backing(r).status for k, r in rows.items()}
            check(got == {"shipped": cb.BACKED, "order": cb.UNBACKED, "later": cb.UNBACKED},
                  f"{what}: carton-backed row kept, order and later rows not", str(got))


def test_vendor_typo_in_a_carton_row() -> None:
    section("a carton row with a typo in its per-carton figure still backs its total")
    with tempfile.TemporaryDirectory() as td:
        path = footwear_like(Path(td) / "typo.xlsx", typo_row=True)
        sheet = cb._Sheet(read_workbook_grids(path)[0])
        rows = rows_of(path)
        v = sheet.backing(rows["shipped"])
        check(v.status == cb.BACKED and len(v.carton_rows) == 5,
              "20 pieces in 1 carton printed as '11 per carton' still counts toward the "
              "total (the 511/533 row-39 case) -- strict-only dropped five real lines",
              v.reason)
        typo = rows["shipped"] - 1
        check(typo not in sheet.strong,
              "but the typo row is WEAK: it can help back a total, never vouch for itself")


def test_weak_rows_are_cartons_only_inside_a_verified_run() -> None:
    section("the weak exception is the typo case only: summaries and strays are not cartons")
    legendz = _sheet(HERE / "Legendz PL0801- 26ctns.xlsx", "PO#1657")
    check({11, 14, 18} <= set(legendz.rejected_weak)
          and not {11, 14, 18} & {r for r, _c in legendz.cartons},
          "Legendz sub-totals R11/R14/R18 (a total + a carton count, the old weak test) are "
          "summaries, not cartons", str(sorted(legendz.rejected_weak)))
    check(legendz.backing(11).status == cb.BACKED,
          "and a sub-total is still BACKED -- by the cartons above it, not by itself",
          legendz.backing(11).reason)
    fw = _sheet(FOOTWEAR, "PO-1624 20139")
    check(39 in fw.rejected_weak and 56 in fw.rejected_weak,
          "footwear's recap rows are not cartons either")

    from openpyxl import Workbook

    with tempfile.TemporaryDirectory() as td:
        path = footwear_like(Path(td) / "typo.xlsx", typo_row=True)
        typo = rows_of(path)["shipped"] - 1
        sheet = cb._Sheet(read_workbook_grids(path)[0])
        check(typo in {r for r, _c in sheet.cartons},
              "the 511/533-style typo row, between verified cartons, is still admitted")

        # The same typo row with a gap on each side: nothing verifies it.
        book = Workbook()
        ws = book.active
        ws.append(["STYLE"] + [""] * 9 + [8, 9, 10, 11, 12, 13, 14])
        for size, count in ((1, 3), (2, 7), (3, 9)):
            row = [""] * 20
            row[10 + size], row[17], row[18], row[19] = 8 * count, 8, count, 8 * count
            ws.append(row)
        ws.append([""] * 20)
        stray = [""] * 20
        stray[1], stray[15], stray[17], stray[18], stray[19] = 26, 20, 11, 1, 20
        ws.append(stray)
        lone = Path(td) / "lone.xlsx"
        book.save(lone)
        sheet = cb._Sheet(read_workbook_grids(lone)[0])
        check(6 in sheet.rejected_weak and 6 not in {r for r, _c in sheet.cartons},
              "a weak row with no verified neighbours on a carton-arithmetic sheet is not a "
              "carton -- it could be a later-shipment row printing a total and a count",
              sheet.rejected_weak.get(6, "admitted"))

        # Two weak rows in ONE run, strong rows the majority: each has the other as
        # a non-strong neighbour, so neither is admitted -- whichever is judged first.
        book = Workbook()
        ws = book.active
        ws.append(["STYLE"] + [""] * 9 + [8, 9, 10, 11, 12, 13, 14])
        for size, qty, per, count in ((1, 24, 8, 3), (2, 56, 8, 7), (3, 30, 11, 1),
                                      (4, 45, 11, 1), (5, 64, 8, 8)):
            row = [""] * 20
            row[1] = 30 + size
            row[10 + size], row[17], row[18], row[19] = qty, per, count, qty
            ws.append(row)
        pair = Path(td) / "pair.xlsx"
        book.save(pair)
        sheet = cb._Sheet(read_workbook_grids(pair)[0])
        check(sheet.strong == {2, 3, 6} and {4, 5} <= set(sheet.rejected_weak),
              "two weak rows sharing a run are BOTH refused -- the verdict does not depend "
              "on which is judged first", f"strong={sorted(sheet.strong)} "
              f"refused={sorted(sheet.rejected_weak)}")


COVERING = HERE / "SD Actual Packing Covering ^N1720^J 1721.pdf"
CARTON_PDF = HERE / "SD Actual Packing ^N1720^J 1721.pdf"
INPROTEX = next(HERE.glob("0626*Invoice_Packing.xlsx"))
LEGENDZ = HERE / "Legendz PL0801- 26ctns.xlsx"

#: The covering PDF's five recap rows, as extracted: 25 lines, 1,669 units.
#: (PO, style, colour, {size: qty}, the carton-detail PDF rows that back it.)
RECAPS_1720_1721 = [
    ("1720", "M650022", "NEW INDIGO",
     {"S": 22, "M": 157, "L": 219, "XL": 150, "2XL": 65, "3XL": 4}, (7, 34)),
    ("1721", "W520005", "COCONUT", {"XS": 33, "S": 79, "M": 65, "L": 33}, (41, 49)),
    ("1721", "W520005", "BLACK", {"XS": 17, "S": 51, "M": 66, "L": 52, "XL": 26}, (50, 58)),
    ("1721", "W600001", "COCONUT", {"XS": 38, "S": 92, "M": 104, "L": 57, "XL": 18},
     (60, 74)),
    ("1721", "W600001", "BLACK", {"XS": 19, "S": 71, "M": 110, "L": 82, "XL": 39}, (75, 89)),
]


def _covering_lines() -> list[dict]:
    return [{"po_number": po, "style_number": style, "color": colour, "size": size,
             "quantity": qty, "source_hint": f"SHEET!R{i + 7}"}
            for i, (po, style, colour, sizes, _rows) in enumerate(RECAPS_1720_1721)
            for size, qty in sizes.items()]


def test_a_total_of_subtotals_is_a_summary() -> None:
    section("the walk-up passes through a refused subtotal: PO TOTAL is not a carton")
    sheet = cb._Sheet(cb.pdf_carton_grid(CARTON_PDF))
    check({35, 36, 59, 90, 91} <= set(sheet.rejected_weak),
          "each colour subtotal AND each PO TOTAL under it is a summary -- the walk no "
          "longer stops at the refused subtotal", str(sorted(sheet.rejected_weak)))
    cartons = [r for r, _c in sheet.cartons]
    check(len(cartons) == 77 and cartons[-1] == 92,
          "leaving the 76 cartons plus G.TOTAL (R92). KNOWN: G.TOTAL still counts -- PO "
          "1720's block sits above a run of header rows the walk does not cross (only a "
          "refused summary is passed through). At 1,669 units it backs no recap here -- "
          "but it WOULD back a recap equal to the grand total (a one-recap document) "
          "on its own, even if the cartons above did not sum to it.",
          f"{len(cartons)} rows, last {cartons[-1]}")


def test_cartons_in_another_attachment_back_a_recap() -> None:
    section("1720/1721: the covering PDF's recaps are backed by the carton-detail PDF")
    lines = _covering_lines()
    result = cb.check_lines(lines, COVERING, evidence=[(CARTON_PDF, CARTON_PDF.name)])
    statuses = [ln["carton_backing"]["status"] for ln in result.kept]
    check(len(result.kept) == 25 and set(statuses) == {cb.BACKED} and not result.dropped,
          "all 25 lines verify against the other attachment's cartons",
          f"{statuses.count(cb.BACKED)} of {len(result.kept)}")
    by_recap = {(ln["style_number"], ln["color"]): ln["carton_backing"] for ln in result.kept}
    for _po, style, colour, _sizes, (first, last) in RECAPS_1720_1721:
        backing = by_recap[(style, colour)]
        check(backing["carton_rows"] == list(range(first, last + 1))
              and backing["carton_source"] == CARTON_PDF.name
              and backing["reason"].count("carton rows") == 1,
              f"{style} {colour}: ONE run, rows {first}-{last}, and no total row with it",
              backing["reason"][:110])


def test_other_pos_cartons_never_back_a_recap() -> None:
    section("negative control: Inprotex's and Legendz's cartons back no 1720/1721 recap")
    lines = _covering_lines()
    result = cb.check_lines(lines, COVERING, evidence=[(INPROTEX, INPROTEX.name),
                                                       (LEGENDZ, LEGENDZ.name)])
    check(all(ln["carton_backing"]["status"] == cb.UNVERIFIABLE for ln in result.kept)
          and len(result.kept) == 25 and not result.dropped,
          "the same email's workbooks for OTHER POs (215 carton rows) back none of the 25 "
          "-- and failing to back is not grounds to drop: all 25 kept, flagged")
    both = cb.check_lines(_covering_lines(), COVERING, evidence=[
        (INPROTEX, INPROTEX.name), (LEGENDZ, LEGENDZ.name), (CARTON_PDF, CARTON_PDF.name)])
    check(all(ln["carton_backing"].get("carton_source") == CARTON_PDF.name
              for ln in both.kept),
          "with all three present, every backing cites the carton-detail PDF alone")

    sheet = cb._Sheet(cb.pdf_carton_grid(CARTON_PDF))
    cartons = dict(sheet.cartons)
    straddle: dict[str, float] = {}
    for row in (34, 41):   # PO 1720's last carton, PO 1721's first: neighbours in the list
        for label, value in cartons[row].items():
            straddle[label] = straddle.get(label, 0.0) + value
    check(cb._runs_summing_to(sheet, straddle) == [],
          "a run never crosses a size header: PO 1720's last carton + PO 1721's first "
          "back nothing, though they sit side by side in the carton list", str(straddle))

    with tempfile.TemporaryDirectory() as td:
        copy = Path(td) / "copy.pdf"
        copy.write_bytes(CARTON_PDF.read_bytes())
        same = cb.check_lines(_covering_lines(), CARTON_PDF, evidence=[(copy, copy.name)])
    check(all(ln["carton_backing"]["status"] == cb.UNVERIFIABLE for ln in same.kept),
          "an attachment with the PRIMARY's own bytes is not independent evidence")


def test_check_lines_drops_keeps_and_flags() -> None:
    section("check_lines: UNBACKED leaves loudly, UNVERIFIABLE stays flagged")
    lines = [
        {"style_number": "20138", "color": "PAT", "size": "10", "quantity": 62,
         "source_hint": "PO-1624 20138!R39"},
        {"style_number": "20138", "color": "PAT", "size": "10", "quantity": 13,
         "source_hint": "PO-1624 20138!R41"},
        {"style_number": "20138", "color": "PAT", "size": "10", "quantity": 62,
         "source_hint": "PAGE1!R7"},
        {"style_number": "20138", "color": "PAT", "size": "10", "quantity": 99,
         "source_hint": "PO-1624 20138!R39"},
        {"style_number": "20138", "color": "PAT", "size": "10", "quantity": 62,
         "source_hint": "PACKING!recap"},
    ]
    result = cb.check_lines([dict(x) for x in lines], FOOTWEAR)
    statuses = [ln["carton_backing"]["status"] for ln in result.kept]
    check(len(result.dropped) == 1 and result.dropped[0]["quantity"] == 13,
          "the unbacked row's line is removed from the shipment", str(result.dropped))
    check(any("NOT SHIPMENT DATA" in w and "row 41" in w for w in result.warnings),
          "and named in a warning, row and lines, so the removal is on the record",
          (result.warnings or ["none"])[0][:90])
    check(statuses == [cb.BACKED, cb.UNVERIFIABLE, cb.UNVERIFIABLE, cb.UNVERIFIABLE],
          "backed kept; missing sheet, wrong-row quantity and row-less hint all KEPT "
          "and flagged unverifiable", str(statuses))
    check("does not hold this line's quantity" in result.kept[2]["carton_backing"]["reason"],
          "a hint pointing at a row without the line's figure is not trusted",
          result.kept[2]["carton_backing"]["reason"][:80])

    with tempfile.TemporaryDirectory() as td:
        pdf_like = Path(td) / "covering.pdf"
        pdf_like.write_bytes(b"%PDF-1.4\n")
        pdf = cb.check_lines([dict(lines[0])], pdf_like)
    check(not pdf.dropped and pdf.kept[0]["carton_backing"]["status"] == cb.UNVERIFIABLE,
          "a PDF has no grid: every line kept, flagged unverifiable -- never dropped "
          "(Paula: keep and flag)", pdf.kept[0]["carton_backing"]["reason"][:70])


def test_inprotex_deterministic_lines_are_all_backed() -> None:
    section("Inprotex, the hand-verified parser: 77 lines row-hinted, 77 backed")
    import document_parsers as dp

    inprotex = next(HERE.glob("0626*Invoice_Packing.xlsx"))
    parsed = dp.parse_packing_slip(inprotex)
    check(parsed.parser == "inprotex-deterministic" and len(parsed.lines) == 77,
          "the deterministic route still produces its 77 verified lines",
          f"{parsed.parser}, {len(parsed.lines)}")
    check(all("!R" in ln["source_hint"] and "recap" not in ln["source_hint"]
              for ln in parsed.lines),
          "every line names its recap ROW -- a row-less hint flagged all 77 for review",
          parsed.lines[0]["source_hint"])
    result = cb.check_lines([dict(ln) for ln in parsed.lines], inprotex)
    statuses = {ln["carton_backing"]["status"] for ln in result.kept}
    check(not result.dropped and statuses == {cb.BACKED},
          "all 77 carton-backed, none dropped, none flagged",
          f"dropped={len(result.dropped)} statuses={statuses}")
    sheet = _sheet(inprotex, "PACKING")
    check(273 not in sheet.header_rows,
          "carton #4 (row 273: `4 | 7 | 9 | 10`, ascending valid sizes) is NOT taken for a "
          "size header -- doing so re-labelled the row beneath it and failed recap row 277")


def test_the_quantity_must_be_under_the_lines_own_size() -> None:
    section("a figure elsewhere on the row, or a multi-row hint, is checked properly")

    def one(size, qty, hint):
        return cb.check_lines([{"style_number": "20139", "color": "DKF", "size": size,
                                "quantity": qty, "source_hint": hint}], FOOTWEAR).kept[0]

    # 20139 R39 prints 8=19, 9=78, 10=148 ...; 148 is on the row, but under size 10.
    wrong_size = one("9", 148, "PO-1624 20139!R39")
    check(wrong_size["carton_backing"]["status"] == cb.UNVERIFIABLE,
          "148 claimed for size 9 is NOT verified just because size 10 holds 148",
          wrong_size["carton_backing"]["reason"][:80])
    formatted = one("9.0", 148, "PO-1624 20139!R39")
    check(formatted["carton_backing"]["status"] == cb.UNVERIFIABLE,
          "nor when the size is written `9.0` -- it normalises like a header cell, instead "
          "of falling through to 'anywhere on the row'")
    right = one("10", 148, "PO-1624 20139!R39")
    check(right["carton_backing"]["status"] == cb.BACKED, "148 for size 10 is")
    both = one("8", 38, "PO-1624 20139!R39, PO-1624 20139!R56")
    check(both["carton_backing"]["status"] == cb.BACKED,
          "a line citing two backed rows passes when its figure is their sum (19 + 19)")
    neither = one("8", 5, "PO-1624 20139!R39, PO-1624 20139!R56")
    check(neither["carton_backing"]["status"] == cb.UNVERIFIABLE,
          "and is flagged when its figure is neither row nor their sum -- multi-row hints "
          "used to skip this check entirely")


def main() -> int:
    print("=" * 78)
    print("CARTON BACKING TESTS -- which quantity rows are shipment data")
    print("=" * 78)
    registered = (
        test_po_1624_real_sheets,
        test_po_1624_sea_shipment_matches_the_hbl,
        test_po_1725_sheets_are_both_backed,
        test_labels_are_never_read,
        test_vendor_typo_in_a_carton_row,
        test_weak_rows_are_cartons_only_inside_a_verified_run,
        test_a_total_of_subtotals_is_a_summary,
        test_cartons_in_another_attachment_back_a_recap,
        test_other_pos_cartons_never_back_a_recap,
        test_check_lines_drops_keeps_and_flags,
        test_inprotex_deterministic_lines_are_all_backed,
        test_the_quantity_must_be_under_the_lines_own_size,
    )
    for fn in registered:
        try:
            fn()
        except Exception:  # noqa: BLE001
            traceback.print_exc()
            _results.append((False, f"{fn.__name__} crashed", ""))
    passed = sum(1 for ok, _n, _d in _results if ok)
    print()
    print(f"{passed}/{len(_results)} checks passed")
    return 0 if passed == len(_results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
