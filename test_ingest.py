"""
Ingest tests -- parser output becoming database rows, idempotently.

Offline: the extractor is stubbed and NetSuite is a mock client, so these run
without an API key or a sandbox. The live end-to-end run against the real corpus
is a separate exercise (see the report in the commit message); what these pin is
the persistence contract, which is the part that must not drift.

    python test_ingest.py
"""

from __future__ import annotations

import datetime as dt
import json
import tempfile
import traceback
from pathlib import Path

from sqlalchemy import func, select, text

import dialect_target as dt_target
import ingest as ing
import schema as sc
from extraction_schema import ParseResult
from netsuite_client import NetSuiteClient, POLine
from schema import (
    attachments,
    audit_log,
    change_candidates,
    message_attachments,
    messages,
    proposed_changes,
    shipment_pos,
    shipment_sources,
    shipments,
    write_attempts,
)

HERE = Path(__file__).resolve().parent
NOW = dt.datetime(2026, 8, 26, 9, 0, 0)
_results: list[tuple[bool, str, str]] = []


def check(ok: bool, name: str, detail: str = "") -> bool:
    _results.append((bool(ok), name, detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  --  {detail}" if detail else ""))
    return bool(ok)


def section(title: str) -> None:
    print()
    print(f"--- {title} " + "-" * max(0, 70 - len(title)))


def fresh_db():
    """
    A built, seeded schema on whichever target is configured.

    SQLite in memory by default; `PO_AGENT_TEST_DB_URL` or `--mssql` points the
    whole file at SQL Server instead. See `dialect_target`.
    """
    return dt_target.fresh_engine()

def counts(engine) -> dict:
    tables = {"messages": messages, "attachments": attachments,
              "message_attachments": message_attachments, "shipments": shipments,
              "shipment_sources": shipment_sources, "shipment_pos": shipment_pos,
              "proposed_changes": proposed_changes, "change_candidates": change_candidates,
              "audit_log": audit_log}
    with engine.connect() as conn:
        return {name: conn.execute(select(func.count()).select_from(t)).scalar()
                for name, t in tables.items()}


def ids(engine, table, column="id") -> set:
    with engine.connect() as conn:
        return set(conn.execute(select(table.c[column])).scalars().all())


def line(po="1662", style="M120246", color="TID", size="S", qty=9,
         conf="high", note="", hint="PACKING!R42"):
    return {"po_number": po, "style_number": style, "color": color, "size": size,
            "quantity": qty, "confidence": conf, "note": note, "source_hint": hint}


def ns_line(line_id="18", style="M120246", color="TID", size="S", qty=12,
            recv=0.0, is_open=True, closed=False):
    return POLine(
        line_id=line_id, item=f"{style} : {style}-{color}-{size}", style_number=style,
        vendor_name="Inprotex", color=color, size=size, quantity=qty, units="Ea",
        expected_receipt_date=dt.date(2026, 7, 6), override_expected_receipt=False,
        updated_receipt_date=None, closed=closed, is_open=is_open,
        quantity_received=recv, quantity_billed=0.0, rate=18.75,
        item_internal_id=f"item-{color}",
    )


class StubExtractor:
    """Stands in for ClaudeExtractor. `model` is read for shipments.extractor_model."""

    model = "claude-opus-5"


def install_stub_parse(monkey: dict, result: ParseResult, classification):
    """
    Replace the two collaborators ingest.py imports at call time.

    They are imported inside `ingest_shipment` precisely so this is possible
    without a DI framework.
    """
    import attachment_classifier
    import document_parsers

    monkey["classify"] = attachment_classifier.classify_attachments
    monkey["parse"] = document_parsers.parse_shipment_email
    attachment_classifier.classify_attachments = (
        lambda paths, extractor=None, **kw: classification)
    document_parsers.parse_shipment_email = (
        lambda paths, extractor=None, cross_check=True, **kw: result
    )


def restore(monkey: dict) -> None:
    import attachment_classifier
    import document_parsers

    attachment_classifier.classify_attachments = monkey["classify"]
    document_parsers.parse_shipment_email = monkey["parse"]


def make_docs(tmp: Path, names=("packing.xlsx",), payload=b"packing-bytes") -> list[Path]:
    out = []
    for index, name in enumerate(names):
        path = tmp / name
        path.write_bytes(payload + bytes([index]))
        out.append(path)
    return out


class FakeClassification:
    """Minimal stand-in for ClassificationResult with real-shaped members."""

    class Item:
        def __init__(self, path, doc_type_value, reason="", excluded_reason=None):
            from attachment_classifier import DocType

            self.path = path
            self.doc_type = DocType(doc_type_value)
            self.reason = reason
            self.unreadable_reason = None
            self.excluded_reason = excluded_reason

    def __init__(self, selected, excluded=()):
        self.selected = list(selected)
        self.excluded = list(excluded)
        self.warnings = []
        self.needs_manual_entry = False

    @property
    def primary(self):
        return self.selected[0] if self.selected else None

    @property
    def cross_checks(self):
        return self.selected[1:]

    def summary(self):
        return "stub"


def msg(graph_id="AAMk-original", **kw):
    return ing.SourceMessage(
        graph_message_id=graph_id, mailbox="shipments@straightdown.com",
        received_at=NOW, subject="SD-219 shipment",
        from_address="exports@inprotex.example", **kw)


# ---------------------------------------------------------------------------


def test_ingest_writes_every_table() -> None:
    section("one ingest, every table, everything the upstream produced")
    engine = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        docs = make_docs(tmp, ("Invoice_Packing.xlsx", "Shipping Advice.pdf"))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list", "PACKING sheet")],
            excluded=[FakeClassification.Item(docs[1], "shipping_advice", "advice",
                                              excluded_reason="not a packing list")])
        parsed = ParseResult(
            lines=[line(size="S", qty=9), line(size="M", qty=50, hint="PACKING!R43"),
                   line(size="XXL", qty=4, conf="low", note="cell smudged",
                        hint="PACKING!R44")],
            ship_info={"etd": "2026/6/27 19:40", "eta": "2026/6/27 16:45"},
            parser="inprotex-deterministic", vendor_name="Inprotex",
            notes=["attachment triage: 1 selected"], warnings=[])
        client = NetSuiteClient(mock_data={"1662": [
            ns_line("18", size="S", qty=12), ns_line("19", size="M", qty=71),
            ns_line("20", size="2X", qty=2)]})

        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            report = ing.ingest_shipment(
                engine, docs, message=msg(), client=client,
                extractor=StubExtractor(), now=NOW)
        finally:
            restore(monkey)

    check(report.created, "the ingest created a shipment", report.summary())
    got = counts(engine)
    for table, expected in (("messages", 1), ("attachments", 2), ("message_attachments", 2),
                            ("shipments", 1), ("shipment_sources", 2), ("shipment_pos", 1),
                            ("proposed_changes", 3), ("audit_log", 1)):
        check(got[table] == expected, f"{table}: {expected} row(s)", str(got[table]))

    with engine.connect() as conn:
        rows = {r.key_size: r for r in conn.execute(select(proposed_changes)).all()}
        ship = conn.execute(select(shipments)).one()
        sources = {r.role for r in conn.execute(select(shipment_sources)).all()}

    # canonical key beside verbatim text -- the pair that makes re-parse safe and
    # still shows Paula what the vendor printed.
    xxl = rows["2x"]
    check(xxl.key_size == "2x" and xxl.src_size_text == "XXL",
          "vendor 'XXL' keys as '2x' while the printed text survives verbatim",
          f"key={xxl.key_size!r} printed={xxl.src_size_text!r}")
    check(rows["s"].key_style == "m120246" and rows["s"].src_style_text == "M120246",
          "style canonicalised and preserved", f"{rows['s'].key_style}/{rows['s'].src_style_text}")
    check(rows["s"].source_hint == "PACKING!R42", "source_hint persisted", rows["s"].source_hint)

    # the five review figures
    check(float(rows["s"].current_quantity) == 12.0
          and float(rows["s"].current_quantity_received) == 0.0
          and float(rows["s"].proposed_quantity) == 9.0
          and rows["s"].ns_line_id == "18",
          "line_balance figures land on the row",
          f"{rows['s'].current_quantity}/{rows['s'].current_quantity_received}"
          f"/{rows['s'].proposed_quantity}/{rows['s'].ns_line_id}")
    with engine.connect() as conn:
        outstanding = conn.execute(text(
            "SELECT outstanding FROM v_review_lines WHERE size_printed = 'S'")).scalar()
    check(float(outstanding) == 12.0, "and outstanding derives in the view", str(outstanding))

    # calibration halves: the claim is stored, the verdict is not yet
    check(xxl.extraction_confidence == "low" and xxl.needs_review == 1,
          "a low-confidence line is marked needs_review",
          f"{xxl.extraction_confidence}/{xxl.needs_review}")
    check(rows["s"].needs_review == 0, "and a high-confidence line is not")
    check(xxl.extraction_note == "cell smudged", "the extractor's note is kept")
    check(all(r.human_verdict is None for r in rows.values()),
          "no human verdict yet -- that is the review UI's job")

    # shipment-level provenance
    # Verbatim as printed, including the time -- normalising on the way in would
    # discard source text. The matcher derives the ISO form for display.
    check(ship.vendor_etd == "2026/6/27 19:40" and ship.vendor_eta == "2026/6/27 16:45",
          "vendor ETD/ETA on the SHIPMENT, verbatim, reference only",
          f"{ship.vendor_etd}/{ship.vendor_eta}")
    check(ship.parser == "inprotex-deterministic" and ship.extractor_model == "claude-opus-5",
          "parser and model recorded for calibration slicing",
          f"{ship.parser}/{ship.extractor_model}")
    check(ship.line_count == 3 and float(ship.unit_total) == 63.0,
          "line count and unit total", f"{ship.line_count}/{ship.unit_total}")
    check(sources == {"PRIMARY", "EXCLUDED"},
          "the advice is recorded as an EXCLUDED source, not parsed", str(sorted(sources)))

    # no date anywhere near a change row
    check(all(r.confirmed_receipt_date is None for r in rows.values()),
          "no confirmed receipt date -- dates never come from a document")


def test_double_ingest_is_a_no_op() -> None:
    section("the same document twice: identical rows, no new ids, no parse")
    engine = fresh_db()
    parse_calls = {"n": 0}
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        docs = make_docs(tmp, ("Invoice_Packing.xlsx",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(lines=[line(size="S"), line(size="M", qty=50)],
                             parser="inprotex-deterministic", vendor_name="Inprotex")
        client = NetSuiteClient(mock_data={"1662": [ns_line("18", size="S"),
                                                    ns_line("19", size="M", qty=71)]})

        import attachment_classifier
        import document_parsers
        keep = (attachment_classifier.classify_attachments,
                document_parsers.parse_shipment_email)

        def counted_parse(paths, extractor=None, cross_check=True, **kw):
            parse_calls["n"] += 1
            return parsed

        attachment_classifier.classify_attachments = (
        lambda paths, extractor=None, **kw: classification)
        document_parsers.parse_shipment_email = counted_parse
        try:
            first = ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
            after_first = counts(engine)
            change_ids = ids(engine, proposed_changes)

            # Same message, same content.
            second = ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
            # And the forwarded case: NEW message id, SAME bytes.
            third = ing.ingest_shipment(
                engine, docs, message=msg("AAMk-forwarded",
                                          forwarded_by="paula@straightdown.com"),
                client=client, now=NOW)
        finally:
            (attachment_classifier.classify_attachments,
             document_parsers.parse_shipment_email) = keep

    check(first.created, "first ingest created the shipment")
    check(not second.created and second.shipment_id == first.shipment_id,
          "second ingest is a no-op returning the same shipment", second.reason[:60])
    check(not third.created and third.shipment_id == first.shipment_id,
          "and a RE-FORWARD with a new message id is too -- content is the axis",
          third.reason[:60])
    check(parse_calls["n"] == 1,
          "the extractor ran ONCE for three ingests -- dedup precedes parsing",
          f"{parse_calls['n']} parse call(s)")

    after = counts(engine)
    for table in ("shipments", "shipment_pos", "proposed_changes", "shipment_sources",
                  "attachments"):
        check(after[table] == after_first[table], f"{table} count unchanged",
              f"{after_first[table]} -> {after[table]}")
    check(ids(engine, proposed_changes) == change_ids,
          "and no new proposed_changes ids were minted")
    check(after["messages"] == 2,
          "the forwarded message IS recorded (provenance), it just starts no shipment",
          str(after["messages"]))
    check(after["audit_log"] == 3,
          "every ingest attempt is audited, including the two skips", str(after["audit_log"]))


def test_po_key_is_canonical() -> None:
    section("one PO written two ways on one document is ONE PO")
    from netsuite_client import po_number_key

    check(po_number_key("1624") == "1624", "bare digits")
    check(po_number_key("PO0001624") == "1624", "a padded tranId collapses to the same key",
          po_number_key("PO0001624"))
    check(po_number_key("PO#1624") == po_number_key(" 1624 ") == po_number_key("PO NO : 1624"),
          "every rendering seen in the corpus lands on one key")
    check(po_number_key("PO0001624") == po_number_key("1624"),
          "THE DEFECT: these two used to be different dict keys and different DB rows")
    # Unresolvable references stay deterministic without inventing a number.
    check(po_number_key("#1720, 1721") == po_number_key("#1720, 1721"),
          "a two-PO reference is still deterministic")
    check(po_number_key("#1720, 1721") not in ("1720", "1721"),
          "but is NOT collapsed onto either one", po_number_key("#1720, 1721"))
    check(po_number_key("") == po_number_key(None) == "", "and empty stays empty")

    # The footwear extraction really did return '1624' for two sheets and
    # 'PO0001624' for the third, and which sheet got which varied between runs.
    # Both splits must produce the same database.
    def run(renderings):
        engine = fresh_db()
        reads = []

        class CountingClient(NetSuiteClient):
            def get_purchase_order(self, po_number):
                reads.append(po_number)
                return super().get_purchase_order(po_number)

        client = CountingClient(mock_data={"1624": [
            ns_line("18", size="S"), ns_line("19", size="M", qty=71),
            ns_line("20", size="L", qty=40)]})
        with tempfile.TemporaryDirectory() as td:
            docs = make_docs(Path(td), ("footwear.xlsx",))
            classification = FakeClassification(
                selected=[FakeClassification.Item(docs[0], "packing_list")])
            parsed = ParseResult(
                lines=[line(po=po, size=size)
                       for po, size in zip(renderings, ("S", "M", "L"))],
                parser="claude-opus-5", vendor_name="Footwear")
            monkey: dict = {}
            install_stub_parse(monkey, parsed, classification)
            try:
                first = ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
                po_rows = _po_rows(engine)
                first_ids = ids(engine, shipment_pos)
                # Same document again: the rows and their ids must not move.
                ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
            finally:
                restore(monkey)
        return first, po_rows, first_ids, ids(engine, shipment_pos), reads, counts(engine)

    mixed = ("1624", "1624", "PO0001624")
    swapped = ("PO0001624", "1624", "1624")

    a_report, a_rows, a_first_ids, a_second_ids, a_reads, a_counts = run(mixed)
    b_report, b_rows, _b_first, _b_second, b_reads, b_counts = run(swapped)

    check(len(a_rows) == 1,
          "THE DEFECT: mixed renderings produce ONE shipment_pos row, not two",
          f"{len(a_rows)} row(s): {[r['po_number_key'] for r in a_rows]}")
    check(a_rows and a_rows[0]["po_number_key"] == "1624",
          "and the stored key is the canonical digits",
          a_rows[0]["po_number_key"] if a_rows else "-")
    check(a_rows and a_rows[0]["ns_tranid"] == "1624",
          "resolution is unaffected (it always normalised)", str(a_rows[0]["ns_tranid"]))
    check(a_counts["proposed_changes"] == 3,
          "all three lines hang off that single parent", str(a_counts["proposed_changes"]))
    check(len(a_reads) == 1, "and the PO was read from NetSuite ONCE, not twice",
          f"{a_reads}")

    check(a_second_ids == a_first_ids,
          "re-ingesting the same document mints no new shipment_pos ids")
    check(len(a_second_ids) == 1, "still exactly one row after the second ingest",
          str(len(a_second_ids)))

    # Run-to-run: the same document, the renderings split differently.
    strip = lambda rows: [{k: v for k, v in r.items() if k != "id"} for r in rows]
    check(strip(a_rows) == strip(b_rows),
          "the two runs produce IDENTICAL shipment_pos rows -- idempotency restored",
          f"{strip(a_rows)} vs {strip(b_rows)}")
    check(a_counts == b_counts, "and identical row counts across every table")
    check(a_reads == b_reads, "and the same NetSuite reads", f"{a_reads} vs {b_reads}")

    # The verbatim renderings are not lost -- they are what the reviewer sees --
    # and BOTH are kept, because picking one of them was itself nondeterministic.
    check(a_rows and a_rows[0]["po_number_printed"] == "1624 / PO0001624",
          "both verbatim renderings are stored, sorted, alongside the key",
          str(a_rows[0]["po_number_printed"]))


def _po_rows(engine) -> list[dict]:
    with engine.connect() as conn:
        return [
            dict(r._mapping)
            for r in conn.execute(
                select(
                    shipment_pos.c.id,
                    shipment_pos.c.po_number_key,
                    shipment_pos.c.po_number_printed,
                    shipment_pos.c.ns_tranid,
                    shipment_pos.c.resolution_status,
                ).order_by(shipment_pos.c.po_number_key)
            ).all()
        ]


def test_size_composition_persisted() -> None:
    section("a composed size reaches the database with both source axes")
    engine = fresh_db()

    def vendor_line(size, primary, secondary, qty, **kw):
        row = line(po="1725", style="50144", color="NIN", size=size, qty=qty)
        row["size_axis_primary"] = primary
        row["size_axis_secondary"] = secondary
        row["size_composition"] = {
            "method": kw.get("method", "COMPOSED"), "primary": primary,
            "secondary": secondary, "composed": size,
        }
        return row

    # Two inseams of one waist: two NetSuite lines, two changes, no collapsing.
    client = NetSuiteClient(mock_data={"1725": [
        ns_line("1", style="50144", color="NIN", size="30-32", qty=16),
        ns_line("8", style="50144", color="NIN", size="30-34", qty=4),
    ]})
    with tempfile.TemporaryDirectory() as td:
        docs = make_docs(Path(td), ("tainan.xls",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(
            lines=[vendor_line("30-32", "30", "INS 32", 17),
                   vendor_line("30-34", "30", "INS 34", 5),
                   # A rejected composition, which must still land with its axes.
                   vendor_line("31-33", "31", "INS 33", 3,
                               method="COMPOSITION_REJECTED")],
            parser="claude-assisted", vendor_name="Tainan")
        monkey: dict = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
        finally:
            restore(monkey)

    with engine.connect() as conn:
        rows = [dict(r._mapping) for r in conn.execute(
            select(proposed_changes.c.src_size_text,
                   proposed_changes.c.key_size,
                   proposed_changes.c.size_composition_method,
                   proposed_changes.c.src_size_axis_primary,
                   proposed_changes.c.src_size_axis_secondary,
                   proposed_changes.c.ns_line_id,
                   proposed_changes.c.proposed_quantity,
                   proposed_changes.c.state)
            .order_by(proposed_changes.c.src_size_text)).all()]

    check(len(rows) == 3, "three rows persisted, not two merged into one", str(len(rows)))
    by_size = {r["src_size_text"]: r for r in rows}

    for size, line_id, qty in (("30-32", "1", 17), ("30-34", "8", 5)):
        row = by_size.get(size, {})
        check(row.get("size_composition_method") == "COMPOSED",
              f"{size}: recorded as COMPOSED", str(row.get("size_composition_method")))
        check(row.get("src_size_axis_primary") == "30",
              f"{size}: the printed waist is stored verbatim",
              str(row.get("src_size_axis_primary")))
        check(row.get("src_size_axis_secondary") == f"INS {size[-2:]}",
              f"{size}: and the printed inseam LABEL, not just its number",
              str(row.get("src_size_axis_secondary")))
        check(row.get("ns_line_id") == line_id,
              f"{size}: matched NetSuite line {line_id}", str(row.get("ns_line_id")))
        check(row.get("proposed_quantity") == qty,
              f"{size}: with its own quantity, not the two summed", str(row.get("proposed_quantity")))

    check(by_size["30-32"]["key_size"] != by_size["30-34"]["key_size"],
          "the two inseams have DIFFERENT canonical keys -- that is what makes "
          "them two lines rather than one of 22")

    rejected = by_size.get("31-33", {})
    check(rejected.get("size_composition_method") == "COMPOSITION_REJECTED",
          "a rejected composition is persisted as such, not dropped",
          str(rejected.get("size_composition_method")))
    check(rejected.get("src_size_axis_primary") == "31"
          and rejected.get("src_size_axis_secondary") == "INS 33",
          "with both axes kept -- the case where they matter most", str(rejected))
    check(rejected.get("ns_line_id") is None and rejected.get("state") == "NEEDS_ATTENTION",
          "and it matched nothing, so it flags", str(rejected.get("state")))

    # The provenance constraint: claiming a composition without both axes is
    # rejected by the database, not merely discouraged.
    from sqlalchemy.exc import IntegrityError

    def half_composition():
        with engine.begin() as conn:
            conn.execute(proposed_changes.update()
                         .where(proposed_changes.c.src_size_text == "30-32")
                         .values(src_size_axis_secondary=None))

    try:
        half_composition()
        check(False, "DB rejects a composition missing an axis", "the update succeeded")
    except IntegrityError as exc:
        check("composition_needs_both_axes" in str(exc),
              "DB rejects a composition missing an axis", str(exc).splitlines()[0][:80])

    # And the single-axis path writes NULLs, not empty strings -- so "was this
    # composed" stays a yes/no question.
    engine2 = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        docs = make_docs(Path(td), ("inprotex.xlsx",), payload=b"other-bytes")
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(lines=[line(size="S")], parser="inprotex-deterministic",
                             vendor_name="Inprotex")
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            ing.ingest_shipment(engine2, docs, message=msg(), client=NetSuiteClient(
                mock_data={"1662": [ns_line("18", size="S")]}), now=NOW)
        finally:
            restore(monkey)
    with engine2.connect() as conn:
        single = conn.execute(select(
            proposed_changes.c.size_composition_method,
            proposed_changes.c.src_size_axis_primary,
            proposed_changes.c.src_size_axis_secondary)).one()
    check(single.size_composition_method is None,
          "a single-axis vendor writes NULL, not a method", str(single.size_composition_method))
    check(single.src_size_axis_primary is None and single.src_size_axis_secondary is None,
          "and NULL axes -- nothing is invented for the four vendors that compose nothing")


def test_transport_mode_recap_rows() -> None:
    section("two recap rows persist as two rows; one line cannot be double-written")
    from sqlalchemy.exc import IntegrityError

    engine = fresh_db()

    def slip(size, qty, label):
        row = line(po="1624", style="20138", color="PAT", size=size, qty=qty)
        row["recap_label"] = label
        return row

    ns = [ns_line("5", style="20138", color="PAT", size="12", qty=44),
          ns_line("43", style="20138", color="PAT", size="12", qty=3, is_open=False)]
    with tempfile.TemporaryDirectory() as td:
        docs = make_docs(Path(td), ("footwear.xlsx",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(
            lines=[slip("12", 52, "By Sea"), slip("12", 5, "By UPS")],
            parser="claude-assisted", vendor_name="Footwear")
        monkey: dict = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            report = ing.ingest_shipment(
                engine, docs, message=msg(), client=NetSuiteClient(mock_data={"1624": ns}),
                now=NOW)
        finally:
            restore(monkey)

    with engine.connect() as conn:
        rows = [dict(r._mapping) for r in conn.execute(
            select(proposed_changes.c.id,
                   proposed_changes.c.key_size,
                   proposed_changes.c.key_recap_label,
                   proposed_changes.c.src_recap_label,
                   proposed_changes.c.state,
                   proposed_changes.c.ns_line_id,
                   proposed_changes.c.proposed_quantity)
            .order_by(proposed_changes.c.key_recap_label)).all()]
        cands = conn.execute(select(func.count()).select_from(change_candidates)).scalar()

    # THE DEFECT: the un-widened canonical-key index forbade this second row.
    check(len(rows) == 2,
          "TWO rows persist for one style/colour/size -- the widened canonical key",
          f"{len(rows)} row(s)")
    check([r["key_recap_label"] for r in rows] == ["by sea", "by ups"],
          "keyed on the canonical recap label", str([r["key_recap_label"] for r in rows]))
    check([r["src_recap_label"] for r in rows] == ["By Sea", "By UPS"],
          "with the verbatim label kept alongside", str([r["src_recap_label"] for r in rows]))
    check(sorted(float(r["proposed_quantity"]) for r in rows) == [5.0, 52.0],
          "each with its own quantity, not summed to 57",
          str(sorted(float(r["proposed_quantity"]) for r in rows)))
    check(all(r["state"] == "NEEDS_ASSIGNMENT" for r in rows),
          "both in NEEDS_ASSIGNMENT", str({r["state"] for r in rows}))
    check(all(r["ns_line_id"] is None for r in rows), "with no target chosen")
    check(report.states.get("NEEDS_ASSIGNMENT") == 2,
          "and the report counts them", str(report.states))
    check(cands == 4,
          "both candidate lines are recorded against BOTH changes, so either can be "
          "assigned to either", f"{cands} candidate row(s)")

    # THE NEW CONSTRAINT: two changes must not select the same NetSuite line.
    def double_write():
        with engine.begin() as conn:
            for r in rows:
                conn.execute(proposed_changes.update()
                             .where(proposed_changes.c.id == r["id"])
                             .values(ns_line_id="5"))

    try:
        double_write()
        check(False, "the DB refuses two changes selecting ONE line", "the update succeeded")
    except IntegrityError as exc:
        # The two engines name the violation differently, and the ONLY portable
        # part is that it was refused -- so check for either idiom rather than
        # asserting on one engine's wording:
        #   SQLite     reports the offending COLUMNS: "proposed_changes.ns_line_id"
        #   SQL Server reports the INDEX: "ux_proposed_changes_one_line_per_shipment"
        # This was the one failure of eight on the first SQL Server run that
        # concerned BEHAVIOUR rather than syntax, and the behaviour turned out to
        # be right on both: the constraint fired. Only the assertion was
        # dialect-specific.
        detail = str(exc)   # NOT `msg` -- that is this module's envelope helper
        by_columns = ("proposed_changes.shipment_id" in detail
                      and "proposed_changes.ns_line_id" in detail)
        by_index = "ux_proposed_changes_one_line_per_shipment" in detail
        check(by_columns or by_index,
              "the DB refuses two changes selecting ONE line, naming that constraint",
              ("columns" if by_columns else "index" if by_index else "NEITHER")
              + ": " + detail.splitlines()[0][-58:])

    # But assigning them to DIFFERENT lines is exactly what should be allowed.
    with engine.begin() as conn:
        for r, target in zip(rows, ("5", "43")):
            conn.execute(proposed_changes.update()
                         .where(proposed_changes.c.id == r["id"])
                         .values(ns_line_id=target))
    with engine.connect() as conn:
        assigned = sorted(conn.execute(select(proposed_changes.c.ns_line_id)).scalars().all())
    check(assigned == ["43", "5"], "a valid assignment to two distinct lines is accepted",
          str(assigned))

    # And a LATER shipment may target the same line again -- the guard is per
    # shipment, not global, because a PO line legitimately gets updated twice.
    engine2 = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        for n, payload in enumerate((b"first", b"second")):
            sub = Path(td) / f"s{n}"
            sub.mkdir()
            docs = make_docs(sub, ("slip.xlsx",), payload=payload)
            classification = FakeClassification(
                selected=[FakeClassification.Item(docs[0], "packing_list")])
            parsed = ParseResult(lines=[line(size="S")], parser="inprotex-deterministic",
                                 vendor_name="Inprotex")
            monkey = {}
            install_stub_parse(monkey, parsed, classification)
            try:
                ing.ingest_shipment(engine2, docs, message=msg(f"AAMk-{n}"),
                                    client=NetSuiteClient(mock_data={
                                        "1662": [ns_line("18", size="S")]}), now=NOW)
            finally:
                restore(monkey)
    with engine2.connect() as conn:
        again = conn.execute(select(func.count()).select_from(proposed_changes)
                             .where(proposed_changes.c.ns_line_id == "18")).scalar()
    check(again == 2,
          "two SHIPMENTS may both target line 18 -- the guard is per shipment", str(again))


def test_multi_po_document() -> None:
    section("one slip, six POs: one shipment_pos row each")
    engine = fresh_db()
    # The real Inprotex sheet interleaves these six.
    po_numbers = ["1640", "1645", "1650", "1662", "1667", "1704"]
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        docs = make_docs(tmp, ("Invoice_Packing.xlsx",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(
            lines=[line(po=po, size="S", qty=10) for po in po_numbers],
            parser="inprotex-deterministic", vendor_name="Inprotex")
        # Only two of the six resolve, which is the point: the other four must not
        # take the shipment down with them.
        client = NetSuiteClient(mock_data={
            "1662": [ns_line("18", size="S", qty=10)],
            "1667": [ns_line("4", size="S", qty=12)]})
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            report = ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
        finally:
            restore(monkey)

    with engine.connect() as conn:
        pos = {r.po_number_key: r for r in conn.execute(select(shipment_pos)).all()}
        changes = conn.execute(select(proposed_changes.c.state)).scalars().all()
    check(len(pos) == 6, "six shipment_pos rows, one per distinct PO", str(len(pos)))
    check(sorted(pos) == po_numbers, "keyed by the printed number", str(sorted(pos)))
    check(pos["1662"].resolution_status == "RESOLVED"
          and pos["1640"].resolution_status == "UNRESOLVED",
          "resolution state is per PO, not per shipment",
          f"1662={pos['1662'].resolution_status} 1640={pos['1640'].resolution_status}")
    check(len(changes) == 6, "and all six lines persisted regardless", str(len(changes)))
    check(sorted(report.states) == ["NEEDS_ATTENTION", "NO_CHANGE", "PENDING_REVIEW"],
          "with a mixed state distribution", str(report.states))


def test_multi_candidate_line() -> None:
    section("a key matching two open lines: candidates persisted, no target chosen")
    engine = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        docs = make_docs(tmp, ("packing.xlsx",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(
            lines=[line(po="1649", style="A320001", color="WHT", size="ALL", qty=58)],
            parser="claude", vendor_name="Symmetry")
        # The real PO0001649 shape: 50 received 0, 200 received 100, both open.
        client = NetSuiteClient(mock_data={"1649": [
            ns_line("1", style="A320001", color="WHT", size="ALL", qty=50, recv=0.0),
            ns_line("2", style="A320001", color="WHT", size="ALL", qty=200, recv=100.0)]})
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
        finally:
            restore(monkey)

    with engine.connect() as conn:
        change = conn.execute(select(proposed_changes)).one()
        cands = conn.execute(select(change_candidates)
                             .order_by(change_candidates.c.ns_line_id)).all()
    check(change.state == sc.STATE_NEEDS_RESOLUTION,
          "the change is NEEDS_RESOLUTION", change.state)
    check(change.ns_line_id is None, "with NO target line chosen", str(change.ns_line_id))
    check(len(cands) == 2, "both candidates persisted", str(len(cands)))
    check([float(c.quantity) for c in cands] == [50.0, 200.0],
          "each with its own quantity -- never summed to 250",
          str([float(c.quantity) for c in cands]))
    check([float(c.quantity_received) for c in cands] == [0.0, 100.0],
          "and its own received figure, which is what a human decides on")
    check(all(c.selected == 0 for c in cands), "nothing pre-selected -- the tool does not pick")
    check(change.current_quantity is None,
          "no line's quantity was adopted as 'current'", str(change.current_quantity))

    # The candidate payload must not carry Paula's manual field.
    with engine.connect() as conn:
        columns = {c["name"] for c in __import__("sqlalchemy").inspect(engine)
                   .get_columns("change_candidates")}
    check(not [c for c in columns if "repspark" in c.lower()],
          "and no repspark column exists to have carried it")


def test_audit_and_state_guard() -> None:
    section("audit trail, and the seeded transition table as runtime authority")
    engine = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        docs = make_docs(tmp, ("packing.xlsx",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(lines=[line()], parser="claude", vendor_name="Inprotex")
        client = NetSuiteClient(mock_data={"1662": [ns_line("18")]})
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            report = ing.ingest_shipment(engine, docs, message=msg(), client=client,
                                         actor="system", now=NOW)
        finally:
            restore(monkey)

    with engine.connect() as conn:
        entries = conn.execute(select(audit_log)).all()
    check(len(entries) == 1, "one audit row for the ingest", str(len(entries)))
    entry = entries[0]
    check(entry.workflow == "PACKING_SLIP" and entry.actor_kind == "SYSTEM",
          "attributed to the packing-slip workflow and the system actor",
          f"{entry.workflow}/{entry.actor_kind}")
    check(entry.event == "SHIPMENT_INGESTED", "with a named event", entry.event)
    check(entry.shipment_id == report.shipment_id and entry.message_id is not None,
          "linked to both the shipment and the message")
    check("states" in (entry.detail_json or ""), "and carrying the state distribution",
          (entry.detail_json or "")[:60])

    # The guard reads change_state_transitions, so the seeded table is load-bearing
    # rather than decoration -- carry-over (a) from the schema review.
    with engine.connect() as conn:
        legal = sc.legal_transitions(conn)
    check((sc.STATE_INSERT, sc.STATE_PENDING_REVIEW) in legal,
          "the insert transition the ingest used is in the table")
    with engine.begin() as conn:
        conn.execute(sc.change_state_transitions.delete().where(
            sc.change_state_transitions.c.from_state == sc.STATE_INSERT).where(
            sc.change_state_transitions.c.to_state == sc.STATE_PENDING_REVIEW))
    try:
        with engine.connect() as conn:
            sc.assert_transition(conn, sc.STATE_INSERT, sc.STATE_PENDING_REVIEW)
        check(False, "removing the row from the TABLE makes the guard refuse",
              "no exception -- the guard is not reading the table")
    except sc.IllegalTransition:
        check(True, "removing the row from the TABLE makes the guard refuse")


def test_colour_resolution_end_to_end() -> None:
    section("colour resolution through the ingest, and what it costs")

    class LiveishClient:
        """
        Duck-types the four things `ingest` asks of a live client, with counters.

        Not a NetSuiteClient subclass on purpose: `is_mock` is the switch ingest
        uses to decide whether to read at all, and faking it on a real client would
        put a mock into paths that refuse mock input.
        """

        is_mock = False

        def __init__(self, lines_by_tranid, colour_names):
            self.lines_by_tranid = lines_by_tranid
            self.colour_names = colour_names
            self.last_lookup_strategy = "stub"
            self.colour_reads = 0

        def resolve_po_internal_id(self, tranid):
            if tranid not in self.lines_by_tranid:
                from netsuite_client import NetSuiteError

                raise NetSuiteError(f"no such PO {tranid}")
            return f"internal-{tranid}"

        def get_purchase_order(self, tranid, **kwargs):
            return self.lines_by_tranid.get(tranid, [])

        def get_item_colour_name(self, item_internal_id, cache=None):
            key = str(item_internal_id)
            if cache is not None and key in cache:
                return cache[key]
            self.colour_reads += 1
            value = self.colour_names.get(key)
            if cache is not None:
                cache[key] = value
            return value

    def po_line(line_id, colour, size, qty, style="M650022"):
        line = ns_line(line_id=line_id, style=style, color=colour, size=size, qty=qty)
        line.item_internal_id = f"item-{colour}"
        return line

    def ingest_one(printed_colour, po_lines, colour_names, style="M650022", db=None):
        engine = db or fresh_db()
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            docs = make_docs(tmp, (f"packing-{printed_colour}.xlsx",),
                             payload=printed_colour.encode())
            classification = FakeClassification(
                selected=[FakeClassification.Item(docs[0], "packing_list")])
            parsed = ParseResult(
                lines=[{"po_number": "1720", "style_number": style,
                        "color": printed_colour, "size": "M", "quantity": 110,
                        "confidence": "high", "note": "", "source_hint": "P1!R3"}],
                parser="claude", vendor_name="Symmetry")
            # Keyed by tranId, because ingest now transforms the printed number
            # before it asks (change 8). A stub keyed by "1720" would fail here,
            # which is the point.
            client = LiveishClient({"PO0001720": po_lines}, colour_names)
            monkey = {}
            install_stub_parse(monkey, parsed, classification)
            try:
                report = ing.ingest_shipment(
                    engine, docs, message=msg(f"AAMk-{printed_colour}"),
                    client=client, now=NOW)
            finally:
                restore(monkey)
        return engine, report, client

    # A code-printing vendor: every printed colour is already a code on the PO, so
    # NOT ONE item is read. This is the cost model the change was scoped around.
    engine, report, client = ingest_one(
        "MLT", [po_line("14", "MLT", "M", 100), po_line("15", "DKF", "M", 50)],
        {"item-MLT": "Moonlight", "item-DKF": "Dark Forest"})
    with engine.connect() as conn:
        row = conn.execute(select(proposed_changes)).one()
    check(row.state == sc.STATE_PENDING_REVIEW and row.ns_line_id == "14",
          "a printed CODE matches through the ingest", f"{row.state}/{row.ns_line_id}")
    check(client.colour_reads == 0,
          "and costs ZERO colour reads -- no lookup was built at all",
          str(client.colour_reads))
    check(report.colour_reads == 0, "the report agrees", str(report.colour_reads))

    # A name-printing vendor: the lookup is built, the name resolves, and the
    # printed text is preserved verbatim beside the canonical key.
    engine, report, client = ingest_one(
        "NEW INDIGO", [po_line("2", "NIN", "M", 155), po_line("3", "MLT", "M", 20)],
        {"item-NIN": "New Indigo", "item-MLT": "Moonlight"})
    with engine.connect() as conn:
        row = conn.execute(select(proposed_changes)).one()
    check(row.state == sc.STATE_PENDING_REVIEW and row.ns_line_id == "2",
          "a printed NAME resolves to the code's line", f"{row.state}/{row.ns_line_id}")
    check(row.src_color_text == "NEW INDIGO" and row.key_color == "new indigo",
          "printed text preserved verbatim, canonical key alongside",
          f"{row.src_color_text!r}/{row.key_color!r}")
    check(client.colour_reads == 2, "one read per distinct colour on the PO",
          str(client.colour_reads))
    check(report.colour_names.get("1720", {}).get("nin") == "New Indigo",
          "and the report records what NetSuite called it",
          str(report.colour_names))

    # An unresolvable colour still flags, and the read attempt is recorded.
    engine, report, client = ingest_one(
        "DFK", [po_line("14", "DKF", "M", 100)], {"item-DKF": "Dark Forest"})
    with engine.connect() as conn:
        row = conn.execute(select(proposed_changes)).one()
    check(row.state == sc.STATE_NEEDS_ATTENTION and row.ns_line_id is None,
          "a colour matching neither a code nor a name flags", row.state)
    check("no NetSuite line" in (row.attention_reason or ""),
          "with the no-match reason", (row.attention_reason or "")[:60])
    check(client.colour_reads == 1,
          "having tried the name path once (the code path missed)",
          str(client.colour_reads))


def test_tranid_resolution() -> None:
    section("printed PO number -> tranId, in the pipeline")
    from netsuite_client import PONumberUnresolvable, po_tranid

    # The rule, against the values whose tranIds are known.
    for printed, expected in (("1662", "PO0001662"), ("1720", "PO0001720"),
                              ("1721", "PO0001721"), ("1657", "PO0001657"),
                              ("7", "PO0000007"), ("1777", "PO0001777")):
        check(po_tranid(printed) == expected, f"{printed!r} -> {expected}", po_tranid(printed))

    # Every rendering seen across the eight real documents.
    for printed in ("PO#1662", "PO NO : 1662", "PO NO  :1662", "PO NO. : 1662",
                    "  1662  ", "PO1662", "1662"):
        check(po_tranid(printed) == "PO0001662",
              f"real-document rendering {printed!r} resolves", po_tranid(printed))

    # Idempotent: applying it to a tranId returns the tranId.
    check(po_tranid("PO0001662") == "PO0001662", "already a tranId -> unchanged")
    check(po_tranid(po_tranid("1662")) == "PO0001662", "and applying it twice is safe")

    # Defined outcomes, not crashes.
    for bad, why in ((("PO NO :"), "no digits"), ("", "empty"), ("   ", "whitespace only")):
        try:
            po_tranid(bad)
            check(False, f"{why} raises PONumberUnresolvable", "no exception")
        except PONumberUnresolvable as exc:
            check(True, f"{why} raises PONumberUnresolvable", str(exc)[:60])

    # Two numbers in one string names two POs. Never pick one.
    try:
        po_tranid("#1720, 1721")
        check(False, "a reference naming TWO POs refuses", "no exception")
    except PONumberUnresolvable as exc:
        check("2 different numbers" in str(exc), "a reference naming TWO POs refuses",
              str(exc)[:80])
        check("wrong order" in str(exc), "and says why picking one would be wrong")
    # ...but a repeated number is one PO, not two.
    check(po_tranid("PO#1662 (1662)") == "PO0001662",
          "the same number twice is still one PO")

    class ResolvingClient:
        """Records what tranId the resolver actually asked for."""

        is_mock = False

        def __init__(self, known):
            self.known = known
            self.asked = []
            self.last_lookup_strategy = None

        def resolve_po_internal_id(self, value):
            from netsuite_client import po_tranid as transform

            tranid = transform(value)
            self.asked.append(tranid)
            if tranid not in self.known:
                raise PONumberUnresolvable(
                    f"PO {value!r} was looked up as tranId {tranid!r} and does not exist",
                    printed=str(value), attempted=tranid)
            self.last_lookup_strategy = "record q= (quoted)"
            return self.known[tranid]

        def get_purchase_order(self, value, **kwargs):
            from netsuite_client import po_tranid as transform

            return [] if transform(value) not in self.known else [ns_line("1")]

        def get_item_colour_name(self, item_internal_id, cache=None):
            return None

    # Resolution through _fetch_po_lines: the printed number is transformed once.
    client = ResolvingClient({"PO0001662": "8489541"})
    lines, resolution = ing._fetch_po_lines(client, ["1662"])
    check(client.asked == ["PO0001662"],
          "the resolver was asked for the tranId, not the printed number", str(client.asked))
    check(resolution["1662"]["status"] == "RESOLVED", "and it resolved",
          resolution["1662"]["status"])
    check(resolution["1662"]["ns_tranid"] == "PO0001662",
          "the derived tranId is recorded", resolution["1662"]["ns_tranid"])
    check(resolution["1662"]["ns_internal_id"] == "8489541", "with the internal id")
    check(resolution["1662"]["strategy"] == "record q= (quoted)",
          "and which q= form worked", str(resolution["1662"]["strategy"]))

    # A PO that does not exist: NOT_FOUND, both strings recorded, no second attempt.
    client = ResolvingClient({"PO0001662": "8489541"})
    lines, resolution = ing._fetch_po_lines(client, ["9999"])
    record = resolution["9999"]
    check(record["status"] == "NOT_FOUND", "an absent PO is NOT_FOUND, not a crash",
          record["status"])
    check(record["ns_tranid"] == "PO0009999",
          "the attempted tranId is still recorded", record["ns_tranid"])
    check("9999" in record["detail"] and "PO0009999" in record["detail"],
          "and the detail carries BOTH the printed value and what was looked up",
          record["detail"][:80])
    check(client.asked == ["PO0009999"],
          "exactly ONE lookup -- no second format was tried", str(client.asked))

    # A malformed reference never reaches NetSuite at all.
    client = ResolvingClient({"PO0001662": "8489541"})
    lines, resolution = ing._fetch_po_lines(client, ["#1720, 1721"])
    check(resolution["#1720, 1721"]["status"] == "NOT_FOUND",
          "a reference naming two POs is NOT_FOUND")
    check(client.asked == [], "and no lookup was attempted", str(client.asked))

    # Per PO, not per shipment: one bad PO does not cost the good ones.
    client = ResolvingClient({"PO0001662": "8489541", "PO0001721": "8669872"})
    lines, resolution = ing._fetch_po_lines(client, ["1662", "9999", "1721"])
    check([resolution[k]["status"] for k in ("1662", "9999", "1721")]
          == ["RESOLVED", "NOT_FOUND", "RESOLVED"],
          "one unresolvable PO leaves the others resolved",
          str([resolution[k]["status"] for k in ("1662", "9999", "1721")]))
    check(sorted(lines) == ["1662", "1721"],
          "and lines come back keyed by the PRINTED number", str(sorted(lines)))

    # The extraction boundary, asserted rather than assumed.
    from netsuite_client import assert_po_reference

    check(assert_po_reference("1662") == "1662", "a PO reference passes the contract")
    try:
        assert_po_reference("")
        check(False, "an empty reference is refused at the boundary", "no exception")
    except PONumberUnresolvable as exc:
        check("extractor" in str(exc),
              "an empty reference is refused, naming the upstream owner", str(exc)[:70])


def test_scope_boundaries() -> None:
    section("scope boundaries the ingest path must not cross")
    engine = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        docs = make_docs(tmp, ("packing.xlsx",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        # An over-ship (500 against 100 ordered) and an orphan line with no
        # NetSuite counterpart, in one shipment.
        parsed = ParseResult(
            lines=[line(size="S", qty=500), line(size="4XL", qty=7)],
            ship_info={"etd": "2026/6/20 08:00", "eta": "2026/7/1 08:00"},
            parser="claude", vendor_name="Inprotex")
        client = NetSuiteClient(mock_data={"1662": [ns_line("18", size="S", qty=100)]})
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
        finally:
            restore(monkey)

    with engine.connect() as conn:
        rows = {r.src_size_text: r for r in conn.execute(select(proposed_changes)).all()}
        ship = conn.execute(select(shipments)).one()

    # Ruling 6: over-shipment is normal and unflagged.
    over = rows["S"]
    check(over.state == sc.STATE_PENDING_REVIEW,
          "500 shipped against 100 ordered is a plain PENDING_REVIEW (ruling 6)", over.state)
    check(over.attention_reason is None, "with no attention reason",
          repr(over.attention_reason))
    for absent in ("PARTIAL_LINE", "OVER_SHIPMENT"):
        check(absent not in (over.attention_reason or ""),
              f"and no {absent} code -- that gate was cancelled")

    # No PO line is ever created: an unmatched line flags instead.
    orphan = rows["4XL"]
    check(orphan.state == sc.STATE_NEEDS_ATTENTION and orphan.ns_line_id is None,
          "an unmatched vendor line flags -- it never becomes a new PO line", orphan.state)
    check("no NetSuite line" in (orphan.attention_reason or ""),
          "and says so", (orphan.attention_reason or "")[:50])

    # Vendor dates reach the shipment and stop there.
    check(ship.vendor_eta == "2026/7/1 08:00",
          "vendor ETA on the shipment, as printed", ship.vendor_eta)
    check(all(r.confirmed_receipt_date is None for r in rows.values()),
          "and no receipt date on any line -- dates come from Paula only")

    # Nothing anywhere mentions the RepSpark field.
    import sqlalchemy as sa

    offenders = [f"{t}.{c['name']}" for t in sa.inspect(engine).get_table_names()
                 for c in sa.inspect(engine).get_columns(t)
                 if "repspark" in c["name"].lower()]
    check(not offenders, "no repspark column in the whole schema", str(offenders))


def test_gaps_are_reported_not_defaulted() -> None:
    section("columns with no upstream producer are reported, not filled in")
    engine = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        docs = make_docs(tmp, ("packing.xlsx",))
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(lines=[line()], parser="claude", vendor_name="Inprotex")
        client = NetSuiteClient(mock_data={"1662": [ns_line("18")]})
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            report = ing.ingest_shipment(engine, docs, message=msg(), client=client, now=NOW)
        finally:
            restore(monkey)

    joined = " | ".join(report.unpopulated)
    for expected in ("agreement_json", "human_verdict"):
        check(expected in joined, f"{expected} reported as unpopulated")

    # These three used to be on the gap list and are now filled, so the assertion
    # flips: reporting a gap that no longer exists is its own kind of stale.
    for closed in ("ns_item_internal_id", "ns_line_is_open", "extractor_prompt_version"):
        check(closed not in joined,
              f"{closed} is no longer reported as a gap -- it has a producer now")

    with engine.connect() as conn:
        row = conn.execute(select(proposed_changes)).one()
        ship = conn.execute(select(shipments)).one()
    check(row.ns_item_internal_id == "item-TID",
          "the matched line's item id persists", str(row.ns_item_internal_id))
    check(row.ns_line_is_open == 1, "and its open state", str(row.ns_line_is_open))
    check(row.ns_line_closed == 0,
          "alongside line_closed, which is NOT its complement", str(row.ns_line_closed))
    check(bool(ship.extractor_prompt_version),
          "and the shipment records which prompt version produced it",
          str(ship.extractor_prompt_version))



# ---------------------------------------------------------------------------
# Multi-batch accumulation (Paula, 2026-09-16)
# ---------------------------------------------------------------------------


def _ship(engine, *, po, slip_lines, ns_lines, graph_id, now=NOW):
    """One full ingest. Distinct bytes per call, or the re-forward dedup fires."""
    with tempfile.TemporaryDirectory() as td:
        docs = make_docs(Path(td), (f"packing-{graph_id}.xlsx",), payload=graph_id.encode())
        classification = FakeClassification(
            selected=[FakeClassification.Item(docs[0], "packing_list")])
        parsed = ParseResult(lines=slip_lines, parser="claude", vendor_name="Inprotex")
        client = NetSuiteClient(mock_data={po: ns_lines})
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        try:
            return ing.ingest_shipment(
                engine, docs, message=msg(graph_id=graph_id), client=client, now=now)
        finally:
            restore(monkey)


def _only_change(engine, shipment_id):
    with engine.connect() as conn:
        return conn.execute(select(proposed_changes)
                            .where(proposed_changes.c.shipment_id == shipment_id)).one()


def _record_write(engine, change, quantity, when, *,
                  state=sc.STATE_WRITTEN, write_status="WRITTEN", outcome="SUCCESS"):
    """Mark a change written the way the write path will, attempt row and all."""
    with engine.begin() as conn:
        conn.execute(proposed_changes.update()
                     .where(proposed_changes.c.id == change.id)
                     .values(state=state, approved_quantity=quantity,
                             quantity_approved_by="paula@straightdown.com",
                             quantity_approved_at=when, quantity_write_status=write_status,
                             updated_at=when))
        conn.execute(write_attempts.insert(), {
            "id": sc.new_id(), "change_id": change.id, "scope": "QUANTITY",
            "attempt_no": 1, "ns_internal_id": "8489541",
            "ns_line_id": change.ns_line_id,
            "payload_json": json.dumps({"quantity": quantity}),
            "idempotency_key": f"{change.id}:QUANTITY:1", "outcome": outcome,
            "http_status": 200 if outcome == "SUCCESS" else 400,
            "error_kind": None if outcome == "SUCCESS" else "TRANSIENT",
            "attempted_at": when,
        })


def test_first_shipment_proposes_the_slip() -> None:
    section("first shipment against an untouched line: the base is zero")
    engine = fresh_db()
    report = _ship(engine, po="1662", graph_id="AAMk-batch-1",
                   slip_lines=[line(qty=128)],
                   ns_lines=[ns_line("18", qty=300, recv=0.0)])
    change = _only_change(engine, report.shipment_id)

    check(float(change.proposed_quantity) == 128.0,
          "proposes the slip's own figure, 128", str(change.proposed_quantity))
    check(change.accumulation_basis == "FIRST_SHIPMENT",
          "basis recorded as FIRST_SHIPMENT", str(change.accumulation_basis))
    check(float(change.accumulation_base_quantity) == 0.0,
          "on a base of zero -- nothing was written before",
          str(change.accumulation_base_quantity))
    check(change.state == sc.STATE_PENDING_REVIEW,
          "and it is an ordinary PENDING_REVIEW", change.state)
    check(float(change.current_quantity) == 300.0,
          "the ORDERED 300 is still recorded as NetSuite's current value, untouched "
          "by the accumulation", str(change.current_quantity))


def test_second_shipment_accumulates() -> None:
    section("second shipment over a WRITTEN first: 128 + 100 = 228, not 100")
    engine = fresh_db()

    first = _ship(engine, po="1662", graph_id="AAMk-batch-1",
                  slip_lines=[line(qty=128)],
                  ns_lines=[ns_line("18", qty=300, recv=0.0)])
    change_1 = _only_change(engine, first.shipment_id)
    _record_write(engine, change_1, 128, dt.datetime(2026, 9, 1, 10, 0, 0))

    # NetSuite now holds what we wrote, and the goods have arrived.
    second = _ship(engine, po="1662", graph_id="AAMk-batch-2",
                   slip_lines=[line(qty=100)],
                   ns_lines=[ns_line("18", qty=128, recv=128.0)],
                   now=dt.datetime(2026, 9, 16, 9, 0, 0))
    change_2 = _only_change(engine, second.shipment_id)

    check(float(change_2.proposed_quantity) == 228.0,
          "proposes 228 -- the written 128 plus this slip's 100",
          str(change_2.proposed_quantity))
    check(float(change_2.proposed_quantity) != 100.0,
          "NOT 100, which is what replace semantics wrote and how 28 units went missing")
    check(change_2.accumulation_basis == "ACCUMULATED",
          "basis ACCUMULATED", str(change_2.accumulation_basis))
    check(float(change_2.accumulation_base_quantity) == 128.0,
          "and the base it added to is recorded, so 228 is legible from the row alone",
          str(change_2.accumulation_base_quantity))
    check(change_2.src_quantity_text == "100",
          "the slip's own figure is still preserved verbatim", str(change_2.src_quantity_text))
    check(change_2.state == sc.STATE_PENDING_REVIEW,
          "an ordinary proposal -- accumulating is not an exception", change_2.state)

    # Provenance: the base came from OUR record of the first write, not off the
    # line. Both read 128 here, so only the citation distinguishes them.
    hist = ing._line_history(engine, ["1662"])[("1662", "18")]
    check(hist.written_quantity == 128.0,
          "the history reports what we wrote", str(hist.written_quantity))
    check(hist.written_change_id == change_1.id,
          "citing the change that wrote it", str(hist.written_change_id))
    check(hist.write_count == 1, "one successful write behind that total", str(hist.write_count))


def test_po_1624_later_shipment_accumulates_onto_the_first() -> None:
    section("PO 1624: the By UPS slip arrives later and ADDS to the By Sea write")
    # Fixture = PO 1624's real figures, 20138 PAT sizes 8-13. The first slip is
    # the carton-backed By Sea row; the second is the By UPS quantities Paula
    # ruled a separate, later shipment -- arriving on their own slip, unlabelled.
    # One NetSuite line per size, the shape the rule is about. (The live PO holds a
    # line per transport mode, which routes the pair to NEEDS_ASSIGNMENT before
    # accumulation is reached; that path is `test_transport_mode_recap_rows`.)
    sizes = ["8", "9", "10", "11", "12", "13"]
    ordered = [10, 40, 75, 85, 55, 25]
    by_sea = [5, 30, 62, 74, 52, 23]
    by_ups = [5, 10, 13, 11, 3, 2]
    line_ids = {s: str(100 + i) for i, s in enumerate(sizes)}

    def lines_for(qtys):
        return [line(po="1624", style="20138", color="PAT", size=s, qty=q,
                     hint="PO-1624 20138!R39") for s, q in zip(sizes, qtys)]

    def ns(qtys, received):
        return [ns_line(line_ids[s], style="20138", color="PAT", size=s, qty=q, recv=r)
                for s, q, r in zip(sizes, qtys, received)]

    def changes(shipment_id):
        with engine.connect() as conn:
            rows = conn.execute(select(proposed_changes).where(
                proposed_changes.c.shipment_id == shipment_id)).all()
        return {r.key_size: r for r in rows}

    engine = fresh_db()
    first = _ship(engine, po="1624", graph_id="AAMk-1624-sea",
                  slip_lines=lines_for(by_sea), ns_lines=ns(ordered, [0.0] * 6))
    first_changes = changes(first.shipment_id)
    check(all(first_changes[s].accumulation_basis == "FIRST_SHIPMENT" for s in sizes),
          "the By Sea slip is each line's first shipment")
    for s, q in zip(sizes, by_sea):
        _record_write(engine, first_changes[s], q, dt.datetime(2026, 5, 23, 10, 0, 0))

    # NetSuite now holds what we wrote, and the sea goods have been received.
    second = _ship(engine, po="1624", graph_id="AAMk-1624-ups",
                   slip_lines=lines_for(by_ups),
                   ns_lines=ns(by_sea, [float(q) for q in by_sea]),
                   now=dt.datetime(2026, 9, 24, 9, 0, 0))
    got = changes(second.shipment_id)

    proposed = [float(got[s].proposed_quantity) for s in sizes]
    check(proposed == [float(o) for o in ordered],
          "each line proposes By Sea + By UPS -- which is exactly the ordered quantity",
          str(proposed))
    check(proposed != [float(q) for q in by_ups],
          "NOT the By UPS figures alone, which is what replace semantics would write")
    check(all(got[s].accumulation_basis == "ACCUMULATED" for s in sizes),
          "basis ACCUMULATED on every line")
    check([float(got[s].accumulation_base_quantity) for s in sizes] == [float(q) for q in by_sea],
          "and the base each added to is the By Sea quantity this tool WROTE",
          str([float(got[s].accumulation_base_quantity) for s in sizes]))
    check(all(got[s].src_quantity_text == str(q) for s, q in zip(sizes, by_ups)),
          "the later slip's own figures are preserved verbatim")
    check(all(got[s].state == sc.STATE_PENDING_REVIEW for s in sizes),
          "an ordinary proposal awaiting Paula's review",
          str(sorted({got[s].state for s in sizes})))


def test_netsuite_disagreeing_with_our_record_is_flagged() -> None:
    section("NetSuite disagrees with what we wrote: flag, do NOT compute")
    engine = fresh_db()

    first = _ship(engine, po="1662", graph_id="AAMk-batch-1",
                  slip_lines=[line(qty=128)],
                  ns_lines=[ns_line("18", qty=300, recv=0.0)])
    change_1 = _only_change(engine, first.shipment_id)
    _record_write(engine, change_1, 128, dt.datetime(2026, 9, 1, 10, 0, 0))

    # Someone edited the line to 150 outside the tool.
    second = _ship(engine, po="1662", graph_id="AAMk-batch-2",
                   slip_lines=[line(qty=100)],
                   ns_lines=[ns_line("18", qty=150, recv=0.0)],
                   now=dt.datetime(2026, 9, 16, 9, 0, 0))
    change_2 = _only_change(engine, second.shipment_id)

    check(change_2.proposed_quantity is None,
          "NOTHING is proposed -- the tool refuses to compute on a base it cannot "
          "vouch for", str(change_2.proposed_quantity))
    check(change_2.accumulation_basis == "DISPUTED",
          "basis DISPUTED", str(change_2.accumulation_basis))
    check(change_2.state == sc.STATE_NEEDS_ATTENTION, "and it goes to Paula", change_2.state)

    reason = change_2.attention_reason or ""
    check("150" in reason, "the reason carries NetSuite's number", reason[:100])
    check("128" in reason, "and what this tool believes it wrote", reason[:100])
    check(change_1.id in reason,
          "citing the change that wrote it, not just the figure")
    check("100" in reason,
          "and the slip's quantity, which is not lost -- only unapplied")

    # THIS is the check that separates the two implementations. An engine basing
    # accumulation on NetSuite's current quantity would return 250 here and look
    # entirely reasonable; one basing it on its own record cannot, because the two
    # sources disagree and that disagreement IS the signal.
    check(change_2.proposed_quantity != 250,
          "and specifically NOT 150 + 100 = 250, which is what reading the base back "
          "out of NetSuite would have produced")


def test_untracked_line_with_receipts_is_flagged() -> None:
    section("no history, goods already arrived: PRE_EXISTING_RECEIPT, not a dispute")
    engine = fresh_db()
    report = _ship(engine, po="1662", graph_id="AAMk-batch-2",
                   slip_lines=[line(qty=100)],
                   ns_lines=[ns_line("18", qty=300, recv=128.0)])
    change = _only_change(engine, report.shipment_id)

    check(change.proposed_quantity is None, "nothing proposed",
          str(change.proposed_quantity))
    check(change.accumulation_basis == "PRE_EXISTING_RECEIPT",
          "basis PRE_EXISTING_RECEIPT", str(change.accumulation_basis))
    check(change.accumulation_basis != "DISPUTED",
          "and specifically NOT DISPUTED -- nothing contradicts anything here, there "
          "is simply no history to add to")
    check(change.state == sc.STATE_NEEDS_ATTENTION,
          "it still needs Paula, because the total has to be confirmed once",
          change.state)

    reason = change.attention_reason or ""
    check("128" in reason, "the reason states what was already received", reason[:110])
    check("100" in reason, "and what this slip adds", reason[:110])
    check("confirm the total" in reason,
          "phrased as a one-time confirmation, not a report of a disagreement",
          reason[:110])
    check("Nothing disagrees" in reason,
          "and says so outright, so a reviewer is not hunting for a conflict that "
          "does not exist", reason[:110])
    check("vendor line has no quantity" not in reason,
          "it does NOT blame the document -- the slip stated 100 perfectly clearly",
          reason[:110])


def test_pre_existing_receipt_retires_itself() -> None:
    section("the day-one case is one-time: confirm it once, then it accumulates")
    engine = fresh_db()

    # First contact with a line that was partly received before the tool existed.
    first = _ship(engine, po="1662", graph_id="AAMk-batch-1",
                  slip_lines=[line(qty=100)],
                  ns_lines=[ns_line("18", qty=300, recv=128.0)])
    change_1 = _only_change(engine, first.shipment_id)
    check(change_1.accumulation_basis == "PRE_EXISTING_RECEIPT",
          "first contact: PRE_EXISTING_RECEIPT", str(change_1.accumulation_basis))

    # Paula confirms the total: the 128 already there plus this slip's 100.
    _record_write(engine, change_1, 228, dt.datetime(2026, 9, 16, 10, 0, 0))

    # A later slip on the same line. NetSuite now holds the confirmed 228.
    second = _ship(engine, po="1662", graph_id="AAMk-batch-2",
                   slip_lines=[line(qty=50)],
                   ns_lines=[ns_line("18", qty=228, recv=228.0)],
                   now=dt.datetime(2026, 10, 1, 9, 0, 0))
    change_2 = _only_change(engine, second.shipment_id)

    check(change_2.accumulation_basis == "ACCUMULATED",
          "the NEXT slip is an ordinary ACCUMULATED -- the flag retired itself",
          str(change_2.accumulation_basis))
    check(change_2.accumulation_basis != "PRE_EXISTING_RECEIPT",
          "it does NOT ask again; the confirmation became the history")
    check(float(change_2.proposed_quantity) == 278.0,
          "and proposes 228 + 50 = 278, with the confirmed total as the base",
          str(change_2.proposed_quantity))
    check(change_2.state == sc.STATE_PENDING_REVIEW,
          "back to an ordinary proposal", change_2.state)


def test_the_two_no_proposal_cases_do_not_share_a_label() -> None:
    section("a genuine contradiction and a missing history read differently")
    # Same outward shape -- nothing proposed, NEEDS_ATTENTION -- and they must
    # still be distinguishable, or the predictable wave trains Paula to skim the
    # word that means alarm.
    engine_a = fresh_db()
    first = _ship(engine_a, po="1662", graph_id="AAMk-batch-1",
                  slip_lines=[line(qty=128)],
                  ns_lines=[ns_line("18", qty=300, recv=0.0)])
    _record_write(engine_a, _only_change(engine_a, first.shipment_id), 128,
                  dt.datetime(2026, 9, 1, 10, 0, 0))
    disputed = _only_change(engine_a, _ship(
        engine_a, po="1662", graph_id="AAMk-batch-2", slip_lines=[line(qty=100)],
        ns_lines=[ns_line("18", qty=150, recv=0.0)],
        now=dt.datetime(2026, 9, 16, 9, 0, 0)).shipment_id)

    engine_b = fresh_db()
    day_one = _only_change(engine_b, _ship(
        engine_b, po="1662", graph_id="AAMk-batch-2", slip_lines=[line(qty=100)],
        ns_lines=[ns_line("18", qty=300, recv=128.0)]).shipment_id)

    check(disputed.proposed_quantity is None and day_one.proposed_quantity is None,
          "both refuse to propose a quantity")
    check(disputed.state == day_one.state == sc.STATE_NEEDS_ATTENTION,
          "and both land in the same state, so the state cannot tell them apart")
    check(disputed.accumulation_basis != day_one.accumulation_basis,
          "but the BASIS does",
          f"{disputed.accumulation_basis} vs {day_one.accumulation_basis}")
    check(disputed.accumulation_basis == "DISPUTED",
          "the contradiction is DISPUTED", str(disputed.accumulation_basis))
    check(day_one.accumulation_basis == "PRE_EXISTING_RECEIPT",
          "the missing history is PRE_EXISTING_RECEIPT", str(day_one.accumulation_basis))
    check("changed the line outside this tool" in (disputed.attention_reason or ""),
          "and only the contradiction accuses anyone of changing anything",
          (disputed.attention_reason or "")[:110])
    check("outside this tool" not in (day_one.attention_reason or ""),
          "the day-one case makes no such accusation, because none happened",
          (day_one.attention_reason or "")[:110])


def test_unwritten_proposals_do_not_move_the_base() -> None:
    section("rejected, pending and failed proposals contribute NOTHING to the base")
    cases = (
        ("never approved", sc.STATE_PENDING_REVIEW, None, None),
        ("rejected outright", sc.STATE_DISCARDED, None, None),
        ("approved but the write FAILED", sc.STATE_WRITE_FAILED, "FAILED", "FAILED"),
    )
    for label, state, write_status, outcome in cases:
        engine = fresh_db()
        first = _ship(engine, po="1662", graph_id="AAMk-batch-1",
                      slip_lines=[line(qty=128)],
                      ns_lines=[ns_line("18", qty=300, recv=0.0)])
        change_1 = _only_change(engine, first.shipment_id)
        if write_status:
            _record_write(engine, change_1, 128, dt.datetime(2026, 9, 1, 10, 0, 0),
                          state=state, write_status=write_status, outcome=outcome)
        elif state != sc.STATE_PENDING_REVIEW:
            with engine.begin() as conn:
                conn.execute(proposed_changes.update()
                             .where(proposed_changes.c.id == change_1.id)
                             .values(state=state, human_verdict="REJECTED",
                                     verdict_by="paula@straightdown.com", verdict_at=NOW))

        # The line still holds 300: nothing was ever written to it.
        second = _ship(engine, po="1662", graph_id="AAMk-batch-2",
                       slip_lines=[line(qty=100)],
                       ns_lines=[ns_line("18", qty=300, recv=0.0)],
                       now=dt.datetime(2026, 9, 16, 9, 0, 0))
        change_2 = _only_change(engine, second.shipment_id)

        check(change_2.accumulation_basis == "FIRST_SHIPMENT",
              f"{label}: still a FIRST_SHIPMENT", str(change_2.accumulation_basis))
        check(float(change_2.accumulation_base_quantity) == 0.0,
              f"{label}: the base stays zero", str(change_2.accumulation_base_quantity))
        check(float(change_2.proposed_quantity) == 100.0,
              f"{label}: proposes the slip alone, 100", str(change_2.proposed_quantity))
        check(float(change_2.proposed_quantity) != 228.0,
              f"{label}: NOT 228 -- an unwritten proposal changed nothing")
        check(float(change_2.proposed_quantity) != 400.0,
              f"{label}: and NOT 300 + 100, which is NetSuite's ordered quantity used "
              "as a base")

        hist = ing._line_history(engine, ["1662"])[("1662", "18")]
        check(hist.written_quantity is None,
              f"{label}: the history records no write", str(hist.written_quantity))
        check(hist.observed_quantity == 300.0,
              f"{label}: but DOES remember seeing the line at 300, which is how an "
              "edit made later would still be caught", str(hist.observed_quantity))


class FlakyClassifierModel:
    """
    A model that answers the SAME question differently on alternate calls.

    Odd-numbered calls say the packing workbook is a size-level packing list;
    even-numbered calls say it is a commercial invoice. That is what the real
    classifier did to PO 1624's clearance workbook on 2026-09-23, reproduced
    offline so the test can prove the pipeline no longer depends on which answer
    it happens to get. The invoice workbook is an invoice on every call.
    """

    model = "flaky-model-1"

    def __init__(self, packing_name: str):
        self.packing_name = packing_name
        self.calls = 0
        self.last_usage = {"input_tokens": 0, "output_tokens": 0}

    def _parse_with_retry(self, schema, system, content):
        self.calls += 1
        self.last_usage = {"input_tokens": self.last_usage["input_tokens"] + 100,
                           "output_tokens": self.last_usage["output_tokens"] + 10}
        verdicts = []
        for block in content[1:]:
            header = block["text"].split("=====")[1]
            if self.packing_name in header and self.calls % 2 == 1:
                verdicts.append({"doc_type": "packing_list", "has_size_breakdown": True,
                                 "reason": "size columns S/M/L with quantities"})
            else:
                verdicts.append({"doc_type": "commercial_invoice", "has_size_breakdown": False,
                                 "reason": "headed COMMERCIAL INVOICE"})
        return schema.model_validate({"verdicts": verdicts})


def _workbook(path: Path, title: str, rows: list) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _ingest_real_classifier(engine, docs, model, parsed):
    """
    Ingest through the REAL classifier and the REAL `parse_shipment_email`.

    Only the packing-slip extraction and the ship-date read are stubbed -- they
    are not what is under test, and they would otherwise need the API. Returns
    the report and how many times the packing slip was parsed.
    """
    import document_parsers

    keep = (document_parsers.parse_packing_slip,
            document_parsers.parse_shipping_info_from_documents)
    parses = []

    def fake_parse(path, extractor=None, **kw):
        import copy

        parses.append(Path(path).name)
        # A copy per call: a primary and a cross-check are separate parses, and
        # `parse_shipment_email` mutates the primary's result in place.
        return copy.deepcopy(parsed)

    document_parsers.parse_packing_slip = fake_parse
    document_parsers.parse_shipping_info_from_documents = lambda paths, extractor=None: ({}, [])
    try:
        report = ing.ingest_shipment(
            engine, docs, message=msg(), client=NetSuiteClient(mock_data={"1662": [
                ns_line("18", size="S"), ns_line("19", size="M", qty=71)]}),
            extractor=model, now=NOW)
    finally:
        (document_parsers.parse_packing_slip,
         document_parsers.parse_shipping_info_from_documents) = keep
    return report, parses


def _classification_snapshot(engine) -> dict:
    with engine.connect() as conn:
        attach = {r.content_sha256: (r.doc_type, r.has_size_breakdown,
                                     r.classifier_rationale, r.classifier_model,
                                     r.classifier_prompt_hash)
                  for r in conn.execute(select(attachments)).all()}
        keys = sorted(
            (r.key_style, r.key_color, r.key_size, r.key_recap_label, r.source_sha256)
            for r in conn.execute(select(proposed_changes)).all())
        roles = sorted((r.content_sha256, r.role)
                       for r in conn.execute(select(shipment_sources)).all())
    return {"attachments": attach, "proposals": keys, "roles": roles}


def test_classification_is_decided_once_and_replayed() -> None:
    section("one classification per ingest, stored, and replayed on re-extraction")
    from attachment_classifier import classifier_prompt_hash

    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        packing = _workbook(tmp / "Clearance Invoice.xlsx", "PO-1662 M120246",
                            [["STYLE", "COLOR", "S", "M", "L"], ["M120246", "TID", 9, 50, 0]])
        invoice = _workbook(tmp / "Commercial Invoice.xlsx", "COMMERCIAL INVOICE",
                            [["COMMERCIAL INVOICE"], ["Total", 59]])
        docs = [packing, invoice]
        parsed = ParseResult(lines=[line(size="S", qty=9), line(size="M", qty=50)],
                             parser="claude-assisted", vendor_name="Test")
        packing_sha = ing.sha256_file(packing)

        # Run 1: a fresh scratch database, nothing on record.
        first_db = fresh_db()
        model = FlakyClassifierModel(packing.name)
        first, parses_1 = _ingest_real_classifier(first_db, docs, model, parsed)
        calls_after_first = model.calls
        snap_1 = _classification_snapshot(first_db)

        # Run 2: the SAME message into a SECOND scratch database that inherits run
        # 1's attachment rows -- which is how `extract_pending --all` re-extracts
        # into a fresh database. Same database would be a dedup no-op that parses
        # nothing and so proves nothing about the decision.
        second_db = fresh_db()
        with first_db.connect() as src, second_db.begin() as dst:
            for row in src.execute(select(attachments)).mappings().all():
                dst.execute(attachments.insert(), dict(row))
        second, parses_2 = _ingest_real_classifier(second_db, docs, model, parsed)
        snap_2 = _classification_snapshot(second_db)

        # Control: the same message into a database with NOTHING on record. The
        # flaky model is now on an even call, so an un-stored classification
        # disagrees. Proves run 2 matched run 1 because of the stored verdict and
        # not because the fake happened to be consistent.
        control_db = fresh_db()
        control, parses_c = _ingest_real_classifier(control_db, docs, model, parsed)
        snap_c = _classification_snapshot(control_db)

    check(first.created and second.created, "both runs created a shipment",
          f"{first.summary()} | {second.summary()}")
    check(calls_after_first == 1,
          "run 1 asked the model ONCE for the whole ingest (was twice: ingest + parse)",
          f"{calls_after_first} call(s)")
    check(parses_1 == [packing.name], "run 1 parsed the workbook its classification selected",
          str(parses_1))
    check(model.calls == calls_after_first + 1,
          "run 2 asked the model NOTHING -- both verdicts came from the attachment rows "
          "(the +1 is the control run)", f"{model.calls} call(s) in total")
    check(snap_1["attachments"] == snap_2["attachments"],
          "identical doc_type, size claim, rationale, model and prompt hash per SHA-256",
          str(sorted(v[0] for v in snap_2["attachments"].values())))
    check(snap_1["proposals"] == snap_2["proposals"] and len(snap_1["proposals"]) == 2,
          "identical proposal keys AND identical source_sha256 across both runs",
          str(snap_2["proposals"]))
    check({p[4] for p in snap_1["proposals"]} == {packing_sha},
          "every proposal names the packing workbook as its source", packing_sha[:12])
    check(snap_1["roles"] == snap_2["roles"]
          and (packing_sha, "PRIMARY") in snap_1["roles"],
          "the RECORDED primary is the document that was PARSED", str(snap_1["roles"]))
    stored = snap_1["attachments"][packing_sha]
    check(stored[0] == "PACKING_LIST" and stored[1] is True
          and stored[2] == "size columns S/M/L with quantities"
          and stored[3] == "flaky-model-1" and stored[4] == classifier_prompt_hash(),
          "the verdict is stored with its rationale, model and prompt hash", str(stored))
    check(snap_c["proposals"] == [] and parses_c == [],
          "control: WITHOUT the stored verdict the flaky model excludes the workbook -- "
          "the stored verdict, not luck, is what made run 2 agree",
          f"{len(snap_c['proposals'])} proposal(s), parsed {parses_c}")


def _flaky_sheet_model():
    """
    A real `ClaudeExtractor` whose model answers the per-SHEET question differently
    on alternate calls.

    The workbook's COMMERCIAL INVOICE sheet comes back a size-level packing list on
    odd-numbered sheet calls and an invoice on even ones -- the per-sheet version
    of what the file-level classifier did to PO 1624 on 2026-09-23. The file-level
    answer is steady (the workbook has sizes), so only the sheet choice can move.
    Extraction calls return one line per sheet and record WHICH sheet was read.
    """
    from types import SimpleNamespace

    import attachment_classifier as ac
    import claude_extractor as ce
    from extraction_schema import PackingSlipExtraction

    packing = {"doc_type": "packing_list", "has_size_breakdown": True,
               "reason": "size columns S/M with quantities"}
    invoice = {"doc_type": "commercial_invoice", "has_size_breakdown": False,
               "reason": "headed COMMERCIAL INVOICE"}

    class FlakySheetModel(ce.ClaudeExtractor):
        def __init__(self):
            super().__init__(client=SimpleNamespace(), model="flaky-sheet-model-1")
            self.sheet_calls = 0
            self.extracted: list[str] = []
            self.last_usage = {"input_tokens": 0, "output_tokens": 0}

        def _parse_with_retry(self, schema, system, content):
            head = content[0]["text"]
            if schema is ac._ContentVerdicts and "section(s)" in head:
                self.sheet_calls += 1
                verdicts = []
                for block in content[1:]:
                    label = block["text"].split("=====")[1].split(":", 1)[1].strip()
                    flip = "INVOICE" in label and self.sheet_calls % 2 == 1
                    verdicts.append(packing if flip or "INVOICE" not in label else invoice)
                return schema.model_validate({"verdicts": verdicts})
            if schema is ac._ContentVerdicts:
                return schema.model_validate({"verdicts": [packing for _ in content[1:]]})
            sent = " ".join(block["text"] for block in content)
            marker = next(m for m in ("MARK-INVOICE", "MARK-PACKING") if m in sent)
            self.extracted.append(marker)
            return PackingSlipExtraction.model_validate({
                "vendor_name": "Test", "document_summary": marker,
                "unparsed_regions": [], "warnings": [],
                "lines": [{"po_number": "1662", "style_number": "M120246", "color": "TID",
                           "size": "S" if marker == "MARK-PACKING" else "M",
                           "quantity": 9, "confidence": "high", "note": "",
                           "source_hint": f"{marker}!recap"}]})

    return FlakySheetModel()


def _ingest_real_parse(engine, docs, model):
    """
    Ingest through the real classifier, the real parse AND the real extractor sheet
    selection. Only the ship-date read is stubbed.
    """
    import document_parsers

    keep = document_parsers.parse_shipping_info_from_documents
    document_parsers.parse_shipping_info_from_documents = lambda paths, extractor=None: ({}, [])
    try:
        return ing.ingest_shipment(
            engine, docs, message=msg(), client=NetSuiteClient(mock_data={"1662": [
                ns_line("18", size="S"), ns_line("19", size="M", qty=71)]}),
            extractor=model, now=NOW)
    finally:
        document_parsers.parse_shipping_info_from_documents = keep


def test_sheet_choice_is_decided_once_and_replayed() -> None:
    section("per-sheet verdicts are stored and replayed: the same sheets, every ingest")
    from openpyxl import Workbook

    from attachment_classifier import classifier_prompt_hash

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "Clearance Invoice.xlsx"
        book = Workbook()
        book.active.title = "COMMERCIAL INVOICE"
        book.active.append(["MARK-INVOICE", "COMMERCIAL INVOICE"])
        book.active.append(["M120246", "TID", "S", 9, "M", 50])
        packing = book.create_sheet("PO-1662 M120246")
        packing.append(["MARK-PACKING", "STYLE", "COLOR", "S", "M"])
        packing.append(["", "M120246", "TID", 9, 50])
        book.save(path)
        sha = ing.sha256_file(path)

        model = _flaky_sheet_model()
        first_db = fresh_db()
        _ingest_real_parse(first_db, [path], model)
        first = list(model.extracted)
        calls_after_first = model.sheet_calls

        # Run 2: the SAME message into a second database that inherits run 1's
        # attachment rows -- how `extract_pending --all` re-extracts. The same
        # database would be a dedup no-op that parses nothing.
        second_db = fresh_db()
        with first_db.connect() as src, second_db.begin() as dst:
            for row in src.execute(select(attachments)).mappings().all():
                dst.execute(attachments.insert(), dict(row))
        model.extracted.clear()
        _ingest_real_parse(second_db, [path], model)
        second = list(model.extracted)
        calls_after_second = model.sheet_calls

        # Control: nothing on record, so the model's NEXT (flipped) answer decides.
        model.extracted.clear()
        _ingest_real_parse(fresh_db(), [path], model)
        control = list(model.extracted)

        with first_db.connect() as conn:
            stored = json.loads(conn.execute(
                select(attachments.c.sheet_verdicts_json)
                .where(attachments.c.content_sha256 == sha)).scalar() or "[]")

    check(calls_after_first == 1, "run 1 asked the sheet question once", str(calls_after_first))
    check(bool(first) and first == second,
          "run 2 extracted EXACTLY the sheets run 1 did", f"{first} vs {second}")
    check(calls_after_second == calls_after_first,
          "because run 2 asked NO sheet question -- every sheet verdict was replayed",
          f"{calls_after_second - calls_after_first} new call(s)")
    check(sorted(control) != sorted(first),
          "control: without the stored verdicts the flaky model picks DIFFERENT sheets -- "
          "the store, not luck, is what made run 2 agree", f"{first} vs control {control}")
    by_sheet = {v["sheet"]: v for v in stored}
    check(set(by_sheet) == {"COMMERCIAL INVOICE", "PO-1662 M120246"}
          and all(v["model"] == "flaky-sheet-model-1"
                  and v["prompt_hash"] == classifier_prompt_hash() for v in stored)
          and by_sheet["PO-1662 M120246"]["rationale"] == "size columns S/M with quantities",
          "each sheet's verdict is stored on the attachment with its rationale, model and "
          "prompt hash", str(sorted(by_sheet)))


def test_a_sheet_verdict_is_reused_only_when_all_four_keys_match() -> None:
    section("a stored sheet verdict needs the same bytes, sheet, prompt AND model")
    from attachment_classifier import SheetVerdicts, StoredSheetVerdict, merge_sheet_verdicts

    v = StoredSheetVerdict("S1", "packing_list", True, "r", "model-a", "hash-1")
    cache = SheetVerdicts({"sha-1": [v]})
    check(cache.lookup("sha-1", "S1", model="model-a", prompt_hash="hash-1") == v,
          "all four match: reused")
    for what, (sha, sheet, model, prompt) in (
            ("other bytes", ("sha-2", "S1", "model-a", "hash-1")),
            ("other sheet", ("sha-1", "S2", "model-a", "hash-1")),
            ("other model", ("sha-1", "S1", "model-b", "hash-1")),
            ("other prompt", ("sha-1", "S1", "model-a", "hash-2"))):
        check(cache.lookup(sha, sheet, model=model, prompt_hash=prompt) is None,
              f"{what}: asked again")
    newer = StoredSheetVerdict("S1", "commercial_invoice", False, "r2", "model-b", "hash-1")
    merged = merge_sheet_verdicts([v], [newer])
    check(len(merged) == 2, "a verdict under another model is ADDED, not overwritten")
    again = StoredSheetVerdict("S1", "commercial_invoice", False, "r3", "model-a", "hash-1")
    remerged = merge_sheet_verdicts(merged, [again])
    check(len(remerged) == 2 and {x.rationale for x in remerged} == {"r2", "r3"},
          "and a verdict under the same key is replaced")


class InvoiceWithSizesModel:
    """Types every workbook commercial_invoice AND says it carries sizes."""

    model = "steady-model-1"

    def __init__(self):
        self.calls = 0
        self.last_usage = {"input_tokens": 0, "output_tokens": 0}

    def _parse_with_retry(self, schema, system, content):
        self.calls += 1
        return schema.model_validate({"verdicts": [
            {"doc_type": "commercial_invoice", "has_size_breakdown": True,
             "reason": "COMMERCIAL INVOICE sheet plus per-PO packing sheets with sizes"}
            for _ in content[1:]]})


def test_two_admitted_workbooks_with_the_same_sheets_do_not_double_propose() -> None:
    section("PO 1624's shape: two invoice-typed workbooks, same packing sheets, one proposal set")
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        rows = [["STYLE", "COLOR", "S", "M"], ["M120246", "TID", 9, 50]]
        clearance = _workbook(tmp / "Clearance Invoice.xlsx", "PO-1662 M120246", rows)
        payment = _workbook(tmp / "Payment invoice - By Sea.xlsx", "PO-1662 M120246",
                            rows + [["forwarder", "X"]])  # different bytes, same sheet data
        parsed = ParseResult(lines=[line(size="S", qty=9), line(size="M", qty=50)],
                             parser="claude-assisted", vendor_name="Test")
        engine = fresh_db()
        report, parses = _ingest_real_classifier(
            engine, [clearance, payment], InvoiceWithSizesModel(), parsed)
        clearance_sha = ing.sha256_file(clearance)
        payment_sha = ing.sha256_file(payment)
        snap = _classification_snapshot(engine)

    check(report.created, "the shipment was created", report.summary())
    check(sorted(parses) == sorted([clearance.name, payment.name]),
          "BOTH workbooks were admitted by the new gate and read (primary + cross-check)",
          str(parses))
    check(len(snap["proposals"]) == 2,
          "two slip lines -> two proposals, NOT four: the cross-check's lines are compared, "
          "never proposed", f"{len(snap['proposals'])} proposal(s)")
    check({p[4] for p in snap["proposals"]} == {clearance_sha},
          "every proposal is sourced from the ONE primary workbook", clearance_sha[:12])
    check(sorted(snap["roles"]) == sorted([(clearance_sha, "PRIMARY"),
                                           (payment_sha, "CROSS_CHECK")]),
          "one PRIMARY, one CROSS_CHECK recorded", str(snap["roles"]))
    check(all(v[0] == "COMMERCIAL_INVOICE" for v in snap["attachments"].values()),
          "doc_type is still recorded, as metadata", str([v[0] for v in snap["attachments"].values()]))


def test_unbacked_rows_never_become_proposals() -> None:
    section("end to end: PO 1624's real workbook, three quantity rows, ONE proposal")
    footwear = HERE / "FW26 footwear PO-1624 packing sheets (invoice sheet removed).xlsx"
    # What the extractor emitted from 20138 PAT size 10 on 2026-09-23: the By Sea
    # row, the Ordered Qty row and the By UPS row (the order row was never emitted
    # by the model, but the rule must hold if it is).
    parsed = ParseResult(lines=[
        line(po="1624", style="20138", color="PAT", size="10", qty=62,
             hint="PO-1624 20138!R39"),
        line(po="1624", style="20138", color="PAT", size="10", qty=75,
             hint="PO-1624 20138!R40"),
        line(po="1624", style="20138", color="PAT", size="10", qty=13,
             hint="PO-1624 20138!R41"),
    ], parser="claude-assisted", vendor_name="Test")
    engine = fresh_db()
    report, _parses = _ingest_real_classifier(engine, [footwear], InvoiceWithSizesModel(),
                                              parsed)
    with engine.connect() as conn:
        rows = conn.execute(select(proposed_changes)).all()
    check([(r.source_hint, r.src_quantity_text) for r in rows] == [("PO-1624 20138!R39", "62")],
          "only the carton-backed row became a proposal", str([r.source_hint for r in rows]))
    dropped = [w for w in report.parse_warnings if w.startswith("NOT SHIPMENT DATA")]
    check(len(dropped) == 2 and any("row 40" in w for w in dropped)
          and any("row 41" in w for w in dropped),
          "and both unbacked rows are named on the shipment's warnings", str(len(dropped)))


def test_lost_rows_escalate_at_the_shipment_level() -> None:
    section("a shipment that loses more than a third of its lines says so on the SHIPMENT")
    import document_parsers as dp

    removed = [{"line": {}, "reason": "NOT SHIPMENT DATA (no carton backing)"}] * 5
    check(dp.removal_escalation(15, removed) is None,
          "exactly a third is not escalated (the rule is MORE than a third)")
    six = dp.removal_escalation(15, removed + removed[:1]) or ""
    check(six.startswith("SHIPMENT ESCALATION") and "6 of 15" in six, "one more line is")
    check(dp.removal_escalation(0, []) is None, "nothing extracted, nothing to escalate")

    footwear = HERE / "FW26 footwear PO-1624 packing sheets (invoice sheet removed).xlsx"
    parsed = ParseResult(lines=[
        line(po="1624", style="20138", color="PAT", size="10", qty=62,
             hint="PO-1624 20138!R39"),
        line(po="", style="20138", color="PAT", size="11", qty=72,
             hint="PO-1624 20138!R39"),
        line(po="1624", style="20138", color="PAT", size="10", qty=75,
             hint="PO-1624 20138!R40"),
        line(po="1624", style="20138", color="PAT", size="10", qty=13,
             hint="PO-1624 20138!R41"),
    ], parser="claude-assisted", vendor_name="Test")
    engine = fresh_db()
    report, _parses = _ingest_real_classifier(engine, [footwear], InvoiceWithSizesModel(),
                                              parsed)
    with engine.connect() as conn:
        ship = conn.execute(select(shipments)).one()
        events = {r.event: r for r in conn.execute(select(audit_log)).all()}
    warnings = json.loads(ship.parse_warnings_json or "[]")
    head = warnings[0] if warnings else ""
    check(head.startswith("SHIPMENT ESCALATION") and "3 of 4" in head,
          "the escalation HEADS the shipment's stored warnings: 2 unbacked + 1 with no PO "
          "number, of 4 extracted", head[:70])
    check(head.index("2 removed as NOT SHIPMENT DATA") < head.index("3 of 4")
          and "rows: PO-1624 20138!R40, PO-1624 20138!R41" in head
          and "1 removed as having no PO number" in head,
          "and it says WHY before how many: each reason, with the rows it removed",
          head[:160])
    check(bool(ship.doc_needs_review), "and the shipment is marked for review")
    rows_lost = events.get("SHIPMENT_ROWS_LOST")
    detail = json.loads(rows_lost.detail_json) if rows_lost else {}
    check(detail.get("lines_extracted") == 4 and detail.get("lines_removed") == 3
          and {d["reason"] for d in detail.get("removed", [])}
          == {"no PO number", "NOT SHIPMENT DATA (no carton backing)"},
          "and a SHIPMENT_ROWS_LOST audit event carries the count and every removed line",
          str(detail)[:120])
    check(any(w.startswith("LINE NOT PERSISTED") and "20138/PAT/11=72" in w for w in warnings),
          "the no-PO line -- once a report-only note -- is named on the shipment itself")
    check(report.parse_warnings[:1] == warnings[:1], "the report leads with it too")


def test_cross_check_compares_normalised_pos_the_primary_covers() -> None:
    section("cross-check: PO numbers normalised, and only the primary's POs compared")
    import document_parsers as dp

    primary = [line(po="1624", style="20138", color="PAT", size="10", qty=62)]
    padded = [line(po="PO0001624", style="20138", color="PAT", size="10", qty=62)]
    got = dp._compare_line_sets(primary, padded, "copy")
    check(len(got) == 1 and "agrees exactly" in got[0],
          "'PO0001624' and '1624' are the same PO -- this read 'DISAGREES' before", got[0])

    other_po = padded + [line(po="1662", style="M120246", color="TID", size="S", qty=9)]
    got = dp._compare_line_sets(primary, other_po, "mixed")[0]
    check("agrees exactly" in got and "1 line(s) on PO(s) 1662 not compared" in got,
          "a line on a PO the primary does not cover is left out, and SAID to be", got)

    unrelated = [line(po="1657", style="M630018", color="DFK", size="M", qty=148)]
    got = dp._compare_line_sets(primary, unrelated, "Legendz")[0]
    check("NOT COMPARED" in got and "DISAGREES" not in got,
          "a document covering none of the primary's POs is not compared at all", got)

    no_po = [line(po="", style="20138", color="PAT", size="10", qty=62)]
    got = dp._compare_line_sets(no_po + primary, no_po, "copy")[0]
    check("NOT COMPARED" in got and "agrees" not in got,
          "two lines with NO PO number are never paired -- that would be agreement "
          "manufactured from style/colour/size alone", got)

    wrong = [line(po="PO0001624", style="20138", color="PAT", size="10", qty=60)]
    got = dp._compare_line_sets(primary, wrong, "copy")[0]
    check("DISAGREES on 1 key" in got, "a real difference on a shared PO still disagrees", got)


def test_cross_check_removals_are_named() -> None:
    section("carton backing on a cross-check names its removals too")
    import attachment_classifier as ac
    import document_parsers

    footwear = HERE / "FW26 footwear PO-1624 packing sheets (invoice sheet removed).xlsx"
    with tempfile.TemporaryDirectory() as td:
        copy_path = Path(td) / "second copy.xlsx"
        copy_path.write_bytes(footwear.read_bytes())

        def verdict(path):
            return ac.AttachmentClassification(
                path=path, doc_type=ac.DocType.PACKING_LIST, has_size_breakdown=True,
                reason="r", method="filename+content", filename_hint=ac.DocType.PACKING_LIST,
                display_name=path.name)

        classification = ac.ClassificationResult(
            selected=[verdict(footwear), verdict(copy_path)])
        lines = [line(po="1624", style="20138", color="PAT", size="10", qty=62,
                      hint="PO-1624 20138!R39"),
                 line(po="1624", style="20138", color="PAT", size="10", qty=13,
                      hint="PO-1624 20138!R41")]
        keep = (document_parsers.parse_packing_slip,
                document_parsers.parse_shipping_info_from_documents)
        document_parsers.parse_packing_slip = (
            lambda path, extractor=None, **kw: ParseResult(
                lines=[dict(x) for x in lines], parser="stub"))
        document_parsers.parse_shipping_info_from_documents = lambda p, extractor=None: ({}, [])
        try:
            result = document_parsers.parse_shipment_email(
                [footwear, copy_path], extractor=StubExtractor(), cross_check=True,
                classification=classification)
        finally:
            (document_parsers.parse_packing_slip,
             document_parsers.parse_shipping_info_from_documents) = keep
    named = [w for w in result.warnings
             if w.startswith("cross-check second copy.xlsx: NOT SHIPMENT DATA")]
    check(len(named) == 1 and "row 41" in named[0],
          "the cross-check's unbacked row is named, with the document it came from",
          (named or ["none"])[0][:90])
    check(result.extracted_line_count == 2 and len(result.removed_lines) == 1,
          "and only the PRIMARY's removals count toward the shipment's lost rows",
          f"{result.extracted_line_count}/{len(result.removed_lines)}")


def test_admitted_shipping_advice_still_supplies_the_dates() -> None:
    section("a shipping advice WITH sizes is now admitted -- and still read for ETD/ETA")
    import attachment_classifier as ac
    import document_parsers

    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        packing, advice = make_docs(tmp, ("packing.xlsx", "Shipping Advice.xlsx"))

        def verdict(path, doc_type):
            return ac.AttachmentClassification(
                path=path, doc_type=doc_type, has_size_breakdown=True, reason="r",
                method="filename+content", filename_hint=doc_type,
                display_name=path.name)

        classification = ac.ClassificationResult(selected=[
            verdict(packing, ac.DocType.PACKING_LIST),
            verdict(advice, ac.DocType.SHIPPING_ADVICE)])
        check(classification.selected[1].usable_as_shipment_data,
              "premise: the advice passes the new gate, so it sits in SELECTED")

        keep = (document_parsers.parse_packing_slip,
                document_parsers.parse_shipping_info_from_documents)
        date_paths = []

        def ship_info(paths, extractor=None):
            date_paths.extend(Path(p).name for p in paths)
            return {"etd": "2026-05-22"}, []

        document_parsers.parse_packing_slip = (
            lambda path, extractor=None, **kw: ParseResult(lines=[line()], parser="stub"))
        document_parsers.parse_shipping_info_from_documents = ship_info
        try:
            result = document_parsers.parse_shipment_email(
                [packing, advice], extractor=StubExtractor(),
                classification=classification)
        finally:
            (document_parsers.parse_packing_slip,
             document_parsers.parse_shipping_info_from_documents) = keep

    check(date_paths == [advice.name],
          "ETD/ETA read from the ADMITTED shipping advice, not the packing list "
          "(it used to be searched for in `excluded` only)", str(date_paths))
    check(result.ship_info.get("etd") == "2026-05-22", "and the date arrives on the result")


def test_proposals_without_a_primary_are_refused() -> None:
    section("invariant: a proposed change with no source_sha256 raises, and writes nothing")
    engine = fresh_db()
    with tempfile.TemporaryDirectory() as td:
        docs = make_docs(Path(td), ("Clearance Invoice.xlsx",))
        # The shape of the 2026-09-23 defect: lines came back from the parse while
        # the recorded classification selected nothing.
        classification = FakeClassification(selected=[], excluded=[
            FakeClassification.Item(docs[0], "commercial_invoice",
                                    excluded_reason="not a packing list")])
        parsed = ParseResult(lines=[line(size="S", qty=9)], parser="claude-assisted")
        monkey = {}
        install_stub_parse(monkey, parsed, classification)
        raised = None
        try:
            ing.ingest_shipment(engine, docs, message=msg(), client=NetSuiteClient(
                mock_data={"1662": [ns_line("18", size="S")]}), extractor=StubExtractor(),
                now=NOW)
        except ing.IngestInvariantError as exc:
            raised = exc
        finally:
            restore(monkey)

    check(raised is not None and "source_sha256" in str(raised),
          "IngestInvariantError, naming source_sha256", str(raised)[:90])
    got = counts(engine)
    check(got["shipments"] == 0 and got["proposed_changes"] == 0
          and got["shipment_sources"] == 0,
          "and the whole shipment rolled back -- no partial rows",
          f"shipments={got['shipments']} proposals={got['proposed_changes']}")


def test_shipment_with_proposals_and_no_recorded_primary_raises() -> None:
    section("invariant: proposals with no PRIMARY source recorded raises")
    engine = fresh_db()
    with engine.begin() as conn:
        conn.execute(shipments.insert(), {
            "id": "s-1", "origin": "PAULA_DIRECTED", "created_by": "test",
            "created_at": NOW, "doc_needs_review": False, "needs_manual_entry": False})
    raised = None
    try:
        with engine.begin() as conn:
            ing._assert_primary_recorded(conn, "s-1", None)
    except ing.IngestInvariantError as exc:
        raised = exc
    check(raised is not None and "0 PRIMARY" in str(raised),
          "a shipment with no PRIMARY shipment_sources row is refused", str(raised)[:90])


def main() -> int:
    print("=" * 78)
    print("INGEST TESTS -- parser output -> database rows")
    print("=" * 78)
    print()
    print("Offline: stubbed extractor, mock NetSuite. Pins the persistence contract;")
    print("the live corpus run is reported separately.")

    REGISTERED = (
        test_ingest_writes_every_table,
        test_double_ingest_is_a_no_op,
        test_po_key_is_canonical,
        test_size_composition_persisted,
        test_transport_mode_recap_rows,
        test_multi_po_document,
        test_multi_candidate_line,
        test_audit_and_state_guard,
        test_colour_resolution_end_to_end,
        test_tranid_resolution,
        test_scope_boundaries,
        test_gaps_are_reported_not_defaulted,
        test_first_shipment_proposes_the_slip,
        test_second_shipment_accumulates,
        test_po_1624_later_shipment_accumulates_onto_the_first,
        test_netsuite_disagreeing_with_our_record_is_flagged,
        test_untracked_line_with_receipts_is_flagged,
        test_pre_existing_receipt_retires_itself,
        test_the_two_no_proposal_cases_do_not_share_a_label,
        test_unwritten_proposals_do_not_move_the_base,
        test_classification_is_decided_once_and_replayed,
        test_sheet_choice_is_decided_once_and_replayed,
        test_a_sheet_verdict_is_reused_only_when_all_four_keys_match,
        test_two_admitted_workbooks_with_the_same_sheets_do_not_double_propose,
        test_admitted_shipping_advice_still_supplies_the_dates,
        test_unbacked_rows_never_become_proposals,
        test_lost_rows_escalate_at_the_shipment_level,
        test_cross_check_removals_are_named,
        test_cross_check_compares_normalised_pos_the_primary_covers,
        test_proposals_without_a_primary_are_refused,
        test_shipment_with_proposals_and_no_recorded_primary_raises,
    )

    # A test registered twice runs twice and its checks are counted twice. That
    # is how this suite reported 115 for 105 distinct checks until a pytest run,
    # which collects each function once, disagreed with the script (RUNBOOK
    # section 8 lessons 18 and 19). Cheap to assert, so the class cannot recur.
    dupes = sorted({f.__name__ for f in REGISTERED if REGISTERED.count(f) > 1})
    check(not dupes, "no test is registered more than once", str(dupes or "none"))

    for fn in REGISTERED:
        try:
            fn()
        except Exception:  # noqa: BLE001
            print()
            traceback.print_exc()
            _results.append((False, f"{fn.__name__} crashed", ""))

    passed = sum(1 for ok, _n, _d in _results if ok)
    total = len(_results)
    print()
    print("=" * 78)
    print(f"{passed}/{total} checks passed")
    print("=" * 78)
    for ok, name, _d in _results:
        if not ok:
            print(f"  FAILED: {name}")
    return 0 if passed == total else 1


if __name__ == "__main__":
    raise SystemExit(main())
